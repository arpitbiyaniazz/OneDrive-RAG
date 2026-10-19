import io
import logging
from typing import List, Optional

from lxml import etree
from pptx import Presentation
from pptx.shapes.group import GroupShape

from app.processors.base import DocumentElement, ParsedDocument, UnsupportedDocumentError, table_elements

logger = logging.getLogger(__name__)

_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_DGM = "http://schemas.openxmlformats.org/drawingml/2006/diagram"


def parse_pptx(data: bytes) -> ParsedDocument:
    try:
        presentation = Presentation(io.BytesIO(data))
    except Exception as exc:
        logger.info("Unreadable presentation: %s", exc)
        raise UnsupportedDocumentError("Unreadable PowerPoint presentation") from exc

    elements: List[DocumentElement] = []
    warnings: List[str] = []
    for number, slide in enumerate(presentation.slides, start=1):
        try:
            elements.extend(_slide_elements(slide, number))
        except Exception:  # one odd slide must not cost the whole deck
            logger.info("Slide %d could not be read in full; using its plain text", number, exc_info=True)
            text = _all_text(slide._element)
            if text:
                elements.append(DocumentElement(text=text, section=f"Slide {number}", page_number=number, kind="slide"))
            warnings.append(f"Slide {number} was read without its layout")
    return ParsedDocument(elements=elements, file_type="pptx", warnings=warnings)


def _slide_elements(slide, number: int) -> List[DocumentElement]:
    title_shape = _title_shape(slide)
    title = _clean(title_shape.text_frame.text) if title_shape is not None and title_shape.has_text_frame else ""
    section = f"Slide {number}: {title}" if title else f"Slide {number}"
    parts: List[str] = []
    tables: List[DocumentElement] = []
    for shape in slide.shapes:
        if title_shape is not None and shape.shape_id == title_shape.shape_id:
            continue
        _collect(shape, slide, parts, tables, section, number)
    notes = _notes(slide)
    if notes:
        parts.append(f"Speaker notes: {notes}")
    body = "\n".join(p for p in parts if p.strip())
    if not body and tables:
        return tables  # the title is already in every table chunk's section
    text = "\n".join(x for x in (title, body) if x)
    elements = [DocumentElement(text=text, section=section, page_number=number, kind="slide")] if text else []
    return elements + tables


def _title_shape(slide):
    try:
        return slide.shapes.title
    except Exception:
        return None


def _notes(slide) -> str:
    try:
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            return _clean(slide.notes_slide.notes_text_frame.text)
    except Exception:
        pass
    return ""


def _collect(shape, slide, parts: List[str], tables: List[DocumentElement], section: str, number: int) -> None:
    try:
        if isinstance(shape, GroupShape):
            for child in shape.shapes:
                _collect(child, slide, parts, tables, section, number)
            return
        if getattr(shape, "has_table", False):
            rows = [[_clean(cell.text) for cell in row.cells] for row in shape.table.rows]
            rows = [r for r in rows if any(r)]
            if len(rows) >= 2:
                tables.extend(table_elements(rows[0], rows[1:], section, page_number=number))
            elif rows:
                parts.append(" | ".join(rows[0]))
            return
        if getattr(shape, "has_chart", False):
            parts.extend(_chart_text(shape.chart))
            return
        smart_art = _smart_art_text(shape, slide)
        if smart_art:
            parts.append(smart_art)
            return
        if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
            parts.extend(_text_frame_lines(shape.text_frame))
        alt = _alt_text(shape)
        if alt:
            parts.append(f"Image: {alt}")
    except Exception:
        logger.debug("Shape skipped", exc_info=True)
        text = _all_text(shape._element)
        if text:
            parts.append(text)


def _text_frame_lines(frame) -> List[str]:
    lines = []
    paragraphs = list(frame.paragraphs)
    for paragraph in paragraphs:
        text = _clean(paragraph.text)  # includes soft line breaks (\v)
        links = []
        for run in paragraph.runs:
            try:
                address = run.hyperlink.address
            except Exception:
                address = None
            if address and address.startswith(("http://", "https://")) and address not in text:
                links.append(address)
        if links:
            text += " (" + ", ".join(links) + ")"
        if text:
            bullet = paragraph.level or len(paragraphs) > 1
            lines.append(("  " * paragraph.level) + f"- {text}" if bullet else text)
    return lines


def _chart_text(chart) -> List[str]:
    """Chart title plus its data as a small table (categories x series)."""
    lines = []
    try:
        if chart.has_title and chart.chart_title.has_text_frame:
            title = _clean(chart.chart_title.text_frame.text)
            if title:
                lines.append(f"Chart: {title}")
    except Exception:
        pass
    try:
        plot = chart.plots[0]
        categories = [_clean(str(c)) for c in plot.categories]
        series = [(s.name or "Series", list(s.values)) for p in chart.plots for s in p.series]
    except Exception:
        return lines
    if categories and series:
        header = ["Category"] + [_clean(str(name)) for name, _ in series]
        lines.append("| " + " | ".join(header) + " |")
        lines.append("| " + " | ".join("---" for _ in header) + " |")
        for i, category in enumerate(categories[:60]):
            values = [_number(vals[i]) if i < len(vals) else "" for _, vals in series]
            lines.append("| " + " | ".join([category] + values) + " |")
    return lines


def _number(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float):
        return f"{value:.15g}"
    return str(value)


def _smart_art_text(shape, slide) -> Optional[str]:
    """SmartArt keeps its text in a separate diagram-data part referenced by the graphic frame."""
    element = shape._element
    rel_ids = element.findall(f".//{{{_DGM}}}relIds")
    if not rel_ids:
        return None
    rel_id = rel_ids[0].get(f"{{{_R}}}dm")
    try:
        part = slide.part.related_part(rel_id)
        root = etree.fromstring(part.blob)
    except Exception:
        return None
    lines = []
    for paragraph in root.iter(f"{{{_A}}}p"):
        text = _clean("".join(t.text or "" for t in paragraph.iter(f"{{{_A}}}t")))
        if text:
            lines.append(f"- {text}")
    return "\n".join(lines) or None


def _alt_text(shape) -> str:
    try:
        props = shape._element.xpath(".//p:cNvPr")
        return _clean(props[0].get("descr") or "") if props else ""
    except Exception:
        return ""


def _all_text(element) -> str:
    """Fallback: every DrawingML text run under an element, paragraph by paragraph."""
    lines = []
    for paragraph in element.iter(f"{{{_A}}}p"):
        text = _clean("".join(t.text or "" for t in paragraph.iter(f"{{{_A}}}t")))
        if text:
            lines.append(text)
    return "\n".join(lines)


def _clean(text: str) -> str:
    return " ".join((text or "").replace("\v", " ").split())
