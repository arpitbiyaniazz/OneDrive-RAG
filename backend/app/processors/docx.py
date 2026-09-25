import io
import docx
from typing import List
from app.processors.base import BaseDocumentParser, DocumentElement, ParsedDocument


class DOCXParser(BaseDocumentParser):
    """Parses Word .docx documents preserving heading structure and sections."""

    def parse(self, file_bytes: bytes, filename: str) -> ParsedDocument:
        elements: List[DocumentElement] = []
        doc = docx.Document(io.BytesIO(file_bytes))

        current_heading = "Introduction"
        current_paragraphs: List[str] = []

        for p in doc.paragraphs:
            text = p.text.strip()
            if not text:
                continue

            # Check if this paragraph is a Heading
            if p.style and p.style.name.startswith("Heading"):
                if current_paragraphs:
                    elements.append(
                        DocumentElement(
                            text="\n".join(current_paragraphs),
                            section=current_heading,
                            metadata={"section": current_heading, "heading_level": p.style.name}
                        )
                    )
                    current_paragraphs = []
                current_heading = text
            else:
                current_paragraphs.append(text)

        # Flush final section
        if current_paragraphs:
            elements.append(
                DocumentElement(
                    text="\n".join(current_paragraphs),
                    section=current_heading,
                    metadata={"section": current_heading}
                )
            )

        # Also extract table text if present
        for t_idx, table in enumerate(doc.tables):
            table_lines: List[str] = []
            for row in table.rows:
                row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                if row_text:
                    table_lines.append(f"| {row_text} |")
            if table_lines:
                elements.append(
                    DocumentElement(
                        text="\n".join(table_lines),
                        section=f"Table {t_idx + 1}",
                        metadata={"type": "table", "table_index": t_idx + 1}
                    )
                )

        return ParsedDocument(
            elements=elements,
            file_type="docx",
            total_pages_or_slides=1,
            metadata={"filename": filename, "format": "docx"}
        )
