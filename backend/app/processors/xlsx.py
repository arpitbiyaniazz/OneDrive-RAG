"""Excel (.xlsx/.xlsm) and table parsing shared with CSV.

A sheet is split into tables at blank rows; each table gets a detected header (title and "Key:
value" rows above it are kept as text, two-row headers are combined, a table without its own
header continues the previous one). Merged cells are filled, number formats applied, and every
table gets a profile (row count, column totals, distinct values) so aggregate questions do not
depend on which row groups happen to be retrieved.
"""
import io
import logging
import re
import zipfile
from collections import Counter
from datetime import date, datetime, time, timedelta
from typing import Any, Dict, List, Optional, Tuple

import openpyxl
from openpyxl.utils.cell import range_boundaries

from app.processors.base import DocumentElement, ParsedDocument, UnsupportedDocumentError, table_elements

logger = logging.getLogger(__name__)

MAX_ROWS_PER_SHEET = 20000
_PROFILE_MIN_ROWS = 5
_PROFILE_MAX_DISTINCT = 20
_HEADER_SCAN_ROWS = 10
_YEAR_RE = re.compile(r"^(19|20)\d{2}$")
_NUMBER_RE = re.compile(r"^[-+]?[$€£¥]?\s?\d[\d,\s ']*(\.\d+)?\s?(%|[$€£¥])?$")
_FORMULA_WITHOUT_VALUE = re.compile(rb"(<f[^>]*>[^<]*</f>|<f[^>]*/>)\s*(<v\s*/>|<v>\s*</v>)?\s*</c>")
_MERGE_RE = re.compile(rb'<mergeCell ref="([A-Z]+[0-9]+:[A-Z]+[0-9]+)"')
_SHEET_RE = re.compile(rb"<sheet\b[^>]*>")


# --------------------------------------------------------------------------- cell values

def format_cell(value: Any, number_format: Optional[str] = None) -> str:
    """Display text of a cell: dates as ISO, percentages and zero-padded codes as formatted, floats
    without binary artifacts (1149.9999999999998 -> 1150)."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == time(0) else value.isoformat(sep=" ", timespec="minutes")
    if isinstance(value, (date, time)):
        return value.isoformat()
    if isinstance(value, timedelta):
        return _hours(value.total_seconds() / 3600)
    if isinstance(value, (int, float)):
        return _format_number(value, number_format or "General")
    return str(value).strip()


def _hours(hours: float) -> str:
    whole = int(hours)
    minutes = int(round((hours - whole) * 60))
    if minutes == 60:
        whole, minutes = whole + 1, 0
    return f"{whole}:{minutes:02d}"


def _format_number(value: float, fmt: str) -> str:
    section = re.sub(r'"[^"]*"|\\.|\[(?!\$)[^\]]*\]', "", fmt.split(";")[0])  # drop literals and [colour]s
    if "[h]" in fmt.lower():
        return _hours(float(value) * 24)
    core = re.sub(r"[$€£¥_*\s()\-]", "", section).strip()
    if "%" in core and re.fullmatch(r"[0#,.%]+", core):
        return f"{float(value) * 100:.{_decimals(core)}f}%"
    if re.fullmatch(r"0{2,}", core) and float(value).is_integer():
        return str(int(value)).zfill(len(core))
    if not core or core.lower() in ("general", "@") or not re.fullmatch(r"[0#,.]+", core):
        return _plain(value)  # General, fractions, scientific and other formats: the plain number
    decimals = _decimals(section)
    grouped = "#,##" in section or "#,#" in section
    text = f"{float(value):,.{decimals}f}" if grouped else f"{float(value):.{decimals}f}"
    symbol = _currency(fmt)
    return f"{symbol}{text}" if symbol else text


def _decimals(section: str) -> int:
    match = re.search(r"\.([0#?]+)", section)
    return len(match.group(1)) if match else 0


def _currency(fmt: str) -> str:
    match = re.search(r'\[\$([^\]-]+)|"([$€£¥])"|(?<!\\)([$€£¥])', fmt)
    return next((g for g in match.groups() if g), "") if match else ""


def _plain(value: float) -> str:
    if isinstance(value, int) or float(value).is_integer():
        return str(int(value))
    return f"{float(value):.15g}"


def _is_number(text: str) -> bool:
    return bool(text) and bool(_NUMBER_RE.match(text.strip()))


def _to_number(text: str) -> Optional[float]:
    cleaned = re.sub(r"[$€£¥%\s ',]", "", text or "")
    try:
        return float(cleaned) if cleaned else None
    except ValueError:
        return None


# --------------------------------------------------------------------------- table structure

def _texty(cell: str) -> bool:
    """Header-like cell: text, or a year label (2024, 2025 above numeric columns)."""
    return bool(cell) and (not _is_number(cell) or bool(_YEAR_RE.match(cell.strip())))


def _header_like(row: List[str], width: int) -> bool:
    filled = [c for c in row if c]
    if len(filled) < 2 or filled[0].rstrip().endswith(":"):  # "Report: | Open orders" is a preamble line
        return False
    return sum(_texty(c) for c in filled) >= len(filled) * 0.6 and len(filled) > width * 0.5


def find_header(rows: List[List[str]]) -> int:
    """Index of the header row of a table: the first row (of the first ten) that spans most of the
    table and is mostly text; 0 when there is none."""
    width = max((len([c for c in r if c]) for r in rows), default=0)
    for i, row in enumerate(rows[:_HEADER_SCAN_ROWS]):
        if _header_like(row, width):
            return i
    return 0


def _combine_header(top: List[str], second: List[str]) -> Optional[List[str]]:
    """Two header rows ("2025" over "Plan | Actual") as one, when the top row groups columns."""
    filled = [c for c in top if c]
    groups = len(set(filled)) < len(filled) or any(not c for c in top[1:])
    if not groups or not any(second):
        return None
    combined, carry = [], ""
    for i, (a, b) in enumerate(zip(top, second, strict=False)):
        if a:
            carry = a
        elif i:
            a = carry  # a group label covers the blank cells to its right
        combined.append(a if a == b or not b else (f"{a} {b}" if a else b))
    return combined


def _split_blocks(numbered: List[Tuple[int, List[str], List[Any]]]) -> List[List[Tuple[int, List[str], List[Any]]]]:
    """Consecutive non-empty rows; blank rows separate tables."""
    blocks, current, last = [], [], None
    for number, row, raw in numbered:
        if last is not None and number > last + 1 and current:
            blocks.append(current)
            current = []
        current.append((number, row, raw))
        last = number
    if current:
        blocks.append(current)
    return blocks


def _profile(section: str, header: List[str], body: List[List[str]], raw: List[List[Any]]) -> Optional[str]:
    if len(body) < _PROFILE_MIN_ROWS:
        return None
    parts = [f"{section}: {len(body)} data rows. Columns: {', '.join(h for h in header if h)}."]
    for col, name in enumerate(header):
        values = [r[col] for r in body if col < len(r) and r[col]]
        if not values:
            continue
        numbers = []
        for i, r in enumerate(body):
            if col < len(r) and r[col]:
                value = raw[i][col] if i < len(raw) and col < len(raw[i]) and isinstance(raw[i][col], (int, float)) \
                    and not isinstance(raw[i][col], bool) else _to_number(r[col]) if _is_number(r[col]) else None
                if value is not None:
                    numbers.append(float(value))
        if len(numbers) >= len(values) * 0.8 and not all(_YEAR_RE.match(v) for v in values):
            total = sum(numbers)
            parts.append(f"{name}: total {_plain(round(total, 6))}, min {_plain(min(numbers))}, "
                         f"max {_plain(max(numbers))}, average {_plain(round(total / len(numbers), 4))}.")
        else:
            counts = Counter(values)
            if 1 < len(counts) <= _PROFILE_MAX_DISTINCT and len(counts) <= len(values) / 2:  # categories, not IDs
                parts.append(f"{name}: " + ", ".join(f"{v} ({n})" for v, n in counts.most_common()) + ".")
    return " ".join(parts)


def _table(section: str, block, previous_header: Optional[List[str]], warnings: List[str],
           ) -> Tuple[List[DocumentElement], Optional[List[str]]]:
    width = max(len(r) for _, r, _ in block)
    rows = [r + [""] * (width - len(r)) for _, r, _ in block]
    raws = [x + [None] * (width - len(x)) for _, _, x in block]
    numbers = [n for n, _, _ in block]
    elements: List[DocumentElement] = []

    has_header = any(_header_like(r, max(len([c for c in r if c]) for r in rows)) for r in rows[:_HEADER_SCAN_ROWS])
    if not has_header and previous_header and len(previous_header) == width:
        header, start = previous_header, 0  # the table continues after a blank row
    elif not has_header and (len(rows) < 3 or width <= 2 or all(r[0].rstrip().endswith(":") for r in rows if r[0])):
        # Notes and "Key: value" lines above or between tables are text, not a table.
        text = "\n".join(" ".join(c for c in r if c) for r in rows)
        return [DocumentElement(text=text, section=section)], previous_header
    else:
        h = find_header(rows)
        if h:
            preamble = "\n".join(" ".join(c for c in r if c) for r in rows[:h] if any(rows))
            if preamble.strip():
                elements.append(DocumentElement(text=preamble, section=section))
        header, start = rows[h], h + 1
        if start < len(rows) - 1 and _header_like(rows[start], width):
            combined = _combine_header(header, rows[start])
            if combined:
                header, start = combined, start + 1
        header = [c or f"Column {i + 1}" for i, c in enumerate(header)]

    body, body_raw, body_numbers = rows[start:], raws[start:], numbers[start:]
    if not body:
        elements.append(DocumentElement(text=" | ".join(header), section=section))
        return elements, header
    if len(body) > MAX_ROWS_PER_SHEET:
        warnings.append(f"{section}: only the first {MAX_ROWS_PER_SHEET} rows were indexed")
        body, body_raw, body_numbers = body[:MAX_ROWS_PER_SHEET], body_raw[:MAX_ROWS_PER_SHEET], body_numbers[:MAX_ROWS_PER_SHEET]
    profile = _profile(section, header, body, body_raw)
    if profile:
        elements.append(DocumentElement(text=profile, section=section, metadata={"profile": True}))
    groups = table_elements(header, body, section, row_numbers=body_numbers)
    span = f"{body_numbers[0]}-{body_numbers[-1]}"
    for element in groups:
        rows_label = element.metadata.get("rows")
        if rows_label and len(groups) > 1:  # say how much of the table this part covers
            element.metadata["rows"] = f"{rows_label} of {span}"
            element.metadata["location"] = f"{section}, rows {rows_label} of {span}"
    elements.extend(groups)
    return elements, header


def rows_to_elements(rows: List[List[str]], section: str, warnings: List[str],
                     raw: Optional[List[List[Any]]] = None) -> List[DocumentElement]:
    """Tables of a sheet or CSV: 1-based row numbers kept, empty columns dropped, split at blank rows."""
    raw = raw or [[None] * len(r) for r in rows]
    numbered = [(i + 1, r, raw[i]) for i, r in enumerate(rows) if any(r)]
    if not numbered:
        return []
    width = max(len(r) for _, r, _ in numbered)
    keep = [i for i in range(width) if any(i < len(r) and r[i] for _, r, _ in numbered)]
    numbered = [(n, [r[i] if i < len(r) else "" for i in keep], [x[i] if i < len(x) else None for i in keep])
                for n, r, x in numbered]
    elements: List[DocumentElement] = []
    previous: Optional[List[str]] = None
    for block in _split_blocks(numbered):
        block_elements, previous = _table(section, block, previous, warnings)
        elements.extend(block_elements)
    return elements


# --------------------------------------------------------------------------- workbook

def _sheet_info(data: bytes) -> Dict[str, Dict[str, Any]]:
    """Per sheet name: visibility and merged ranges, read straight from the package XML (read-only
    mode exposes neither)."""
    info: Dict[str, Dict[str, Any]] = {}
    try:
        package = zipfile.ZipFile(io.BytesIO(data))
        workbook = package.read("xl/workbook.xml")
        rels = package.read("xl/_rels/workbook.xml.rels").decode("utf-8", "replace")
    except (KeyError, zipfile.BadZipFile):
        return info
    targets = {m.group(1): m.group(2) for m in re.finditer(r'<Relationship\b[^>]*Id="([^"]+)"[^>]*Target="([^"]+)"', rels)}
    targets.update({m.group(2): m.group(1) for m in re.finditer(r'<Relationship\b[^>]*Target="([^"]+)"[^>]*Id="([^"]+)"', rels)})
    for tag in _SHEET_RE.findall(workbook):
        text = tag.decode("utf-8", "replace")
        name = re.search(r'name="([^"]*)"', text)
        rel = re.search(r'r:id="([^"]+)"', text)
        state = re.search(r'state="([^"]+)"', text)
        if not name:
            continue
        sheet_name = _unescape(name.group(1))
        target = targets.get(rel.group(1), "") if rel else ""
        path = target.lstrip("/") if target.startswith("/") else f"xl/{target}"
        info[sheet_name] = {"hidden": bool(state and state.group(1) != "visible"), "merges": _merges(package, path),
                            "formula_without_value": _formulas_without_values(package, path)}
    return info


def _unescape(text: str) -> str:
    return text.replace("&amp;", "&").replace("&lt;", "<").replace("&gt;", ">").replace("&quot;", '"').replace("&apos;", "'")


def _scan(package: zipfile.ZipFile, path: str, pattern: re.Pattern, first_only: bool = False) -> List[bytes]:
    found: List[bytes] = []
    try:
        with package.open(path) as stream:
            tail = b""
            while True:
                chunk = stream.read(1 << 20)
                if not chunk:
                    break
                window = tail + chunk
                found.extend(m.group(0) if not m.groups() else m.group(1) for m in pattern.finditer(window))
                if first_only and found:
                    return found
                tail = window[-256:]
    except KeyError:
        return []
    return list(dict.fromkeys(found))


def _merges(package: zipfile.ZipFile, path: str) -> List[Tuple[int, int, int, int]]:
    ranges = []
    for ref in _scan(package, path, _MERGE_RE):
        try:
            min_col, min_row, max_col, max_row = range_boundaries(ref.decode("ascii"))
            ranges.append((min_row, min_col, max_row, max_col))
        except ValueError:
            continue
    return ranges


def _formulas_without_values(package: zipfile.ZipFile, path: str) -> bool:
    return bool(_scan(package, path, _FORMULA_WITHOUT_VALUE, first_only=True))


def parse_xlsx(data: bytes) -> ParsedDocument:
    try:
        values_wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:
        logger.info("Unreadable spreadsheet: %s", exc)
        raise UnsupportedDocumentError("Unreadable Excel workbook") from exc
    info = _sheet_info(data)
    formulas_wb = None
    if any(sheet.get("formula_without_value") for sheet in info.values()):
        # Files saved by some tools carry no cached formula results; show the formula instead.
        formulas_wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=False)

    elements: List[DocumentElement] = []
    warnings: List[str] = []
    try:
        for ws in values_wb.worksheets:  # chart sheets have no cells
            sheet = info.get(ws.title, {})
            section = f"Sheet: {ws.title}" + (" (hidden)" if sheet.get("hidden") else "")
            ws.reset_dimensions()  # the stored <dimension> can be wrong ("A1"), which would hide every row
            formula_rows = None
            if formulas_wb is not None and sheet.get("formula_without_value") and ws.title in formulas_wb.sheetnames:
                formula_ws = formulas_wb[ws.title]
                formula_ws.reset_dimensions()
                formula_rows = formula_ws.iter_rows(values_only=True)
            rows: List[List[str]] = []
            raw: List[List[Any]] = []
            for cells in ws.iter_rows():
                formulas = next(formula_rows, ()) if formula_rows is not None else ()
                row, row_raw = [], []
                for index, cell in enumerate(cells):
                    value = getattr(cell, "value", None)
                    if value is None and index < len(formulas) and isinstance(formulas[index], str) \
                            and formulas[index].startswith("="):
                        row.append(formulas[index])
                        row_raw.append(None)
                        continue
                    row.append(format_cell(value, getattr(cell, "number_format", None)))
                    row_raw.append(value)
                rows.append(row)
                raw.append(row_raw)
                if len(rows) > MAX_ROWS_PER_SHEET + 50:
                    break
            _fill_merges(rows, raw, sheet.get("merges", []))
            elements.extend(rows_to_elements(rows, section, warnings, raw))
    finally:
        values_wb.close()
        if formulas_wb is not None:
            formulas_wb.close()
    return ParsedDocument(elements=elements, file_type="xlsx", warnings=warnings)


def _fill_merges(rows: List[List[str]], raw: List[List[Any]], merges: List[Tuple[int, int, int, int]]) -> None:
    """Every cell of a merged range shows the range's value (a region label merged over three rows)."""
    for min_row, min_col, max_row, max_col in merges:
        if min_row - 1 >= len(rows) or (max_row - min_row + 1) * (max_col - min_col + 1) > 10000:
            continue
        top = rows[min_row - 1]
        value = top[min_col - 1] if min_col - 1 < len(top) else ""
        value_raw = raw[min_row - 1][min_col - 1] if min_col - 1 < len(raw[min_row - 1]) else None
        if not value:
            continue
        for r in range(min_row - 1, min(max_row, len(rows))):
            row, row_raw = rows[r], raw[r]
            if len(row) < max_col:
                row.extend([""] * (max_col - len(row)))
                row_raw.extend([None] * (max_col - len(row_raw)))
            for c in range(min_col - 1, max_col):
                if not row[c]:
                    row[c], row_raw[c] = value, value_raw
