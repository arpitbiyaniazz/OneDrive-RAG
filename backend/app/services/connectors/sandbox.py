"""In-memory sandbox drives with generated sample documents (DEV_MOCK_* only, never in production)."""
import io
from datetime import UTC, datetime
from typing import Callable, Dict, List, Optional, Tuple

from app.services.connectors.base import DriveItem, NotFoundError

_MODIFIED = datetime(2026, 9, 1, 9, 0, tzinfo=UTC)
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
PPTX = "application/vnd.openxmlformats-officedocument.presentationml.presentation"


def _pdf(pages: List[str]) -> bytes:
    import pymupdf

    doc = pymupdf.open()
    for text in pages:
        doc.new_page().insert_textbox(pymupdf.Rect(50, 60, 545, 800), text, fontsize=11)
    buf = io.BytesIO()
    doc.save(buf)
    doc.close()
    return buf.getvalue()


def _docx(blocks: List[Tuple[str, str]]) -> bytes:
    import docx

    d = docx.Document()
    for kind, text in blocks:
        if kind.startswith("h"):
            d.add_heading(text, level=int(kind[1]))
        else:
            d.add_paragraph(text)
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _xlsx(title: str, rows: List[list]) -> bytes:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = title
    for row in rows:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _pptx(slides: List[Tuple[str, str]]) -> bytes:
    from pptx import Presentation

    prs = Presentation()
    for title, body in slides:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = body
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


# (id, name, folder id, folder path, mime, content factory)
_ONEDRIVE_FILES: List[Tuple[str, str, str, str, str, Callable[[], bytes]]] = [
    ("file_welcome", "Company Overview.txt", "root", "/", "text/plain", lambda: (
        b"Welcome to the Contoso knowledge base. Contoso was founded in 2020 and has three core "
        b"departments: HR, Finance and Engineering.")),
    ("file_hr_leave", "Leave Policy 2026.pdf", "folder_hr", "/HR", "application/pdf", lambda: _pdf([
        "CONTOSO LEAVE POLICY 2026\n\n1. Annual leave: employees are entitled to 18 annual leave days per "
        "calendar year. Requests must be submitted through the HR portal at least two weeks in advance.\n\n"
        "2. Sick leave: employees receive 10 paid sick days annually. A medical certificate is required for "
        "absences longer than 3 consecutive business days.\n\n3. Carry over: up to 5 unused days can be carried "
        "into the first quarter of the following year.",
        "Parental and special leave\n\nPrimary caregivers receive 16 weeks of fully paid parental leave. "
        "Secondary caregivers receive 4 weeks of paid leave. Bereavement leave provides up to 5 paid days.",
    ])),
    ("file_hr_handbook", "Employee Handbook.docx", "folder_hr", "/HR", DOCX, lambda: _docx([
        ("h1", "Leave management"),
        ("p", "All employees record time off in the HR portal. Annual leave accrues at 1.5 days per month."),
        ("h1", "Remote work"),
        ("p", "Employees may work remotely up to 3 days per week with team alignment."),
    ])),
    ("file_fin_budget", "Budget 2026.xlsx", "folder_finance", "/Finance", XLSX, lambda: _xlsx("Budget 2026", [
        ["Department", "Q1 Budget", "Q2 Budget", "Q3 Budget", "Q4 Budget", "Total 2026"],
        ["Engineering", 1200000, 1350000, 1400000, 1500000, 5450000],
        ["Marketing", 450000, 500000, 550000, 600000, 2100000],
        ["Operations", 300000, 310000, 320000, 330000, 1260000],
    ])),
    ("file_fin_expense", "Expense Policy.txt", "folder_finance", "/Finance", "text/plain", lambda: (
        b"Contoso Expense Policy 2026\n1. Meals during business travel are reimbursed up to $75 per day.\n"
        b"2. Flights over 4 hours may be booked in Premium Economy with manager approval.\n"
        b"3. Receipts must be uploaded within 30 days.")),
    ("file_eng_arch", "Microservices Architecture.pptx", "folder_engineering", "/Engineering", PPTX, lambda: _pptx([
        ("Microservices Architecture", "Platform overview"),
        ("Data pipelines", "Ingestion chunks documents by section\nEmbeddings are stored in pgvector\nEvery query is filtered by tenant"),
    ])),
]
_ONEDRIVE_FOLDERS = [("folder_hr", "HR"), ("folder_finance", "Finance"), ("folder_engineering", "Engineering")]

_GDRIVE_FILES: List[Tuple[str, str, str, str, str, Callable[[], bytes]]] = [
    ("gfile_security", "Enterprise AI Security Policy", "gfolder_corporate", "/Corporate", "application/vnd.google-apps.document",
     lambda: _docx([
         ("h1", "Data loss prevention"),
         ("p", "All assistant queries are scrubbed for personal data before prompt construction."),
         ("h1", "Model access"),
         ("p", "Only approved models with zero data retention may process enterprise text."),
     ])),
    ("gfile_costs", "Global Infrastructure Costs 2026", "gfolder_corporate", "/Corporate", "application/vnd.google-apps.spreadsheet",
     lambda: _xlsx("Cloud costs", [
         ["Provider", "Category", "Q1 Actual ($)", "Annual Budget ($)"],
         ["Google Cloud Platform", "Kubernetes & BigQuery", 185000, 780000],
         ["Microsoft Azure", "Azure OpenAI & storage", 145000, 600000],
     ])),
    ("gfile_briefing", "Executive Briefing Q3", "root", "/", "application/vnd.google-apps.presentation", lambda: _pptx([
        ("Executive Briefing Q3", "Multi-cloud document assistant"),
        ("Milestones", "Google Drive and OneDrive connectors\nHybrid search over all document chunks"),
    ])),
    ("gfile_k8s", "Kubernetes Hybrid Cloud Architecture.pdf", "gfolder_tech", "/Engineering Specs", "application/pdf", lambda: _pdf([
        "KUBERNETES HYBRID CLOUD ARCHITECTURE\n\nWorkloads span GKE and AKS with an Istio service mesh. "
        "Recovery point objective: under 60 seconds. Recovery time objective: under 5 minutes.",
    ])),
]
_GDRIVE_FOLDERS = [("gfolder_corporate", "Corporate"), ("gfolder_tech", "Engineering Specs")]


class SandboxConnector:
    """Implements the Connector protocol over a fixed in-memory tree. Tokens are ignored."""

    def __init__(self, drive_type: str, folders, files, url_prefix: str):
        self.drive_type = drive_type
        self._folders = folders
        self._file_specs = files
        self._url_prefix = url_prefix
        self.reset()

    def reset(self) -> None:
        self._content: Dict[str, bytes] = {}
        self._modified: Dict[str, datetime] = {}
        self._items: Dict[str, Tuple[DriveItem, str]] = {}
        for folder_id, name in self._folders:
            self._items[folder_id] = (DriveItem(folder_id, name, True, self.drive_type, "folder", folder_path="/"), "root")
        for file_id, name, parent, path, mime, factory in self._file_specs:
            self._content[file_id] = factory()
            self._items[file_id] = (self._file_item(file_id, name, mime, path), parent)

    def _file_item(self, file_id: str, name: str, mime: str, path: str) -> DriveItem:
        return DriveItem(
            id=file_id, name=name, is_folder=False, drive_type=self.drive_type, mime_type=mime,
            size=len(self._content[file_id]), modified_date=self._modified.get(file_id, _MODIFIED),
            created_date=_MODIFIED, web_url=f"{self._url_prefix}{file_id}", folder_path=path,
        )

    # Test hooks
    def set_content(self, file_id: str, data: bytes, modified: Optional[datetime] = None) -> None:
        item, parent = self._items[file_id]
        self._content[file_id] = data
        self._modified[file_id] = modified or datetime.now(UTC)
        self._items[file_id] = (self._file_item(file_id, item.name, item.mime_type, item.folder_path), parent)

    def remove(self, file_id: str) -> None:
        self._items.pop(file_id, None)
        self._content.pop(file_id, None)

    def add_file(self, file_id: str, name: str, parent_id: str, data: bytes, mime: str = "text/plain") -> None:
        path = "/" if parent_id == "root" else self._items[parent_id][0].path
        self._content[file_id] = data
        self._modified[file_id] = datetime.now(UTC)
        self._items[file_id] = (self._file_item(file_id, name, mime, path), parent_id)

    def rename(self, file_id: str, name: str) -> None:
        item, parent = self._items[file_id]
        self._modified[file_id] = datetime.now(UTC)
        self._items[file_id] = (self._file_item(file_id, name, item.mime_type, item.folder_path), parent)

    # Connector protocol
    async def get_root(self, token: str) -> DriveItem:
        return DriveItem("root", "", True, self.drive_type, "folder", folder_path="/")

    async def list_children(self, token: str, folder_id: str, folder_path: str) -> List[DriveItem]:
        return [item for item, parent in self._items.values() if parent == folder_id]

    async def get_file(self, token: str, file_id: str) -> DriveItem:
        if file_id == "root":
            return await self.get_root(token)
        if file_id not in self._items:
            raise NotFoundError("Sandbox item not found")
        return self._items[file_id][0]

    async def download(self, token: str, item: DriveItem) -> bytes:
        if item.id not in self._content:
            raise NotFoundError("Sandbox item not found")
        return self._content[item.id]


sandbox_onedrive = SandboxConnector("onedrive", _ONEDRIVE_FOLDERS, _ONEDRIVE_FILES, "https://onedrive.live.com/?id=")
sandbox_gdrive = SandboxConnector("google_drive", _GDRIVE_FOLDERS, _GDRIVE_FILES, "https://drive.google.com/file/d/")
