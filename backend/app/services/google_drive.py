import io
import logging
from datetime import datetime, timezone
from typing import Any, Dict, List, Optional
import httpx

from app.core.config import settings
from app.core.telemetry.otel_setup import trace_gdrive_call

logger = logging.getLogger(__name__)

GOOGLE_AUTH_URL = "https://accounts.google.com/o/oauth2/v2/auth"
GOOGLE_TOKEN_URL = "https://oauth2.googleapis.com/token"
GOOGLE_USERINFO_URL = "https://www.googleapis.com/oauth2/v2/userinfo"
GOOGLE_DRIVE_API_BASE = "https://www.googleapis.com/drive/v3"


class GoogleDriveService:
    """
    Production client for Google Drive API v3 and Google OAuth 2.0.
    Supports directory exploration, binary streaming, Google Workspace document export,
    incremental changes/delta tracking, and push notification webhooks.
    """
    def __init__(self):
        self.client_id = settings.GOOGLE_CLIENT_ID
        self.client_secret = settings.GOOGLE_CLIENT_SECRET
        self.scopes = [
            "https://www.googleapis.com/auth/drive.readonly",
            "openid",
            "https://www.googleapis.com/auth/userinfo.email",
            "https://www.googleapis.com/auth/userinfo.profile",
        ]

    def get_auth_url(self, redirect_uri: str, state: str) -> str:
        """Constructs Google OAuth 2.0 authorization URL."""
        scope_str = "%20".join(self.scopes)
        return (
            f"{GOOGLE_AUTH_URL}"
            f"?client_id={self.client_id}"
            f"&response_type=code"
            f"&redirect_uri={redirect_uri}"
            f"&scope={scope_str}"
            f"&state={state}"
            f"&access_type=offline"
            f"&prompt=consent"
        )

    async def exchange_code(self, code: str, redirect_uri: str) -> Dict[str, Any]:
        """Exchanges authorization code for Google access and refresh tokens."""
        data = {
            "client_id": self.client_id,
            "client_secret": self.client_secret,
            "grant_type": "authorization_code",
            "code": code,
            "redirect_uri": redirect_uri,
        }
        async with httpx.AsyncClient() as client:
            res = await client.post(GOOGLE_TOKEN_URL, data=data)
            if res.status_code != 200:
                logger.error(f"Failed to exchange Google OAuth code: {res.text}")
                raise ValueError(f"Google OAuth failed: {res.text}")
            return res.json()

    async def refresh_access_token(self, refresh_token: str) -> Dict[str, Any]:
        """Refreshes an expired Google access token using the refresh token."""
        data = {
            "client_id": self.client_id,
            "client_secret": self.client_secret,
            "grant_type": "refresh_token",
            "refresh_token": refresh_token,
        }
        async with httpx.AsyncClient() as client:
            res = await client.post(GOOGLE_TOKEN_URL, data=data)
            if res.status_code != 200:
                raise ValueError(f"Google token refresh failed: {res.text}")
            return res.json()

    async def get_user_profile(self, access_token: str) -> Dict[str, Any]:
        """Fetches the user profile from Google userinfo."""
        with trace_gdrive_call("get_user_profile", "/userinfo"):
            headers = {"Authorization": f"Bearer {access_token}"}
            async with httpx.AsyncClient() as client:
                res = await client.get(GOOGLE_USERINFO_URL, headers=headers)
                if res.status_code != 200:
                    raise ValueError(f"Failed to fetch Google profile: {res.text}")
                return res.json()

    async def list_drive_items(self, access_token: str, folder_id: str = "root") -> List[Dict[str, Any]]:
        """Lists files and folders within a Google Drive folder."""
        with trace_gdrive_call("list_drive_items", f"/files?q='{folder_id}' in parents"):
            headers = {"Authorization": f"Bearer {access_token}"}
            query = f"'{folder_id}' in parents and trashed = false"
            params = {
                "q": query,
                "fields": "files(id, name, mimeType, size, createdTime, modifiedTime, webViewLink, iconLink)",
                "pageSize": 100,
            }
            async with httpx.AsyncClient() as client:
                res = await client.get(f"{GOOGLE_DRIVE_API_BASE}/files", headers=headers, params=params)
                if res.status_code != 200:
                    raise ValueError(f"Failed to list Google Drive files: {res.text}")
                data = res.json()
                items = data.get("files", [])
                formatted = []
                for item in items:
                    mime = item.get("mimeType", "")
                    is_folder = (mime == "application/vnd.google-apps.folder")
                    formatted.append({
                        "id": item["id"],
                        "name": item["name"],
                        "is_folder": is_folder,
                        "size": int(item.get("size", 0)),
                        "mime_type": mime,
                        "created_date": item.get("createdTime"),
                        "modified_date": item.get("modifiedTime"),
                        "web_url": item.get("webViewLink", ""),
                        "drive_type": "google_drive",
                    })
                return formatted

    async def get_file_metadata(self, access_token: str, file_id: str) -> Dict[str, Any]:
        """Retrieves metadata for a specific Google Drive file."""
        with trace_gdrive_call("get_file_metadata", f"/files/{file_id}"):
            headers = {"Authorization": f"Bearer {access_token}"}
            params = {
                "fields": "id, name, mimeType, size, createdTime, modifiedTime, webViewLink, parents"
            }
            async with httpx.AsyncClient() as client:
                res = await client.get(f"{GOOGLE_DRIVE_API_BASE}/files/{file_id}", headers=headers, params=params)
                if res.status_code != 200:
                    raise ValueError(f"Failed to get Google Drive file metadata: {res.text}")
                return res.json()

    async def download_file_bytes(self, access_token: str, file_id: str, mime_type: str = "") -> bytes:
        """
        Downloads binary content of a file or exports Google Workspace native documents
        (Docs -> DOCX, Sheets -> XLSX, Slides -> PPTX).
        """
        with trace_gdrive_call("download_file_bytes", f"/files/{file_id}/download"):
            headers = {"Authorization": f"Bearer {access_token}"}
            async with httpx.AsyncClient(timeout=60.0) as client:
                # Handle Google Workspace Native Docs
                if mime_type == "application/vnd.google-apps.document":
                    # Export Google Doc as DOCX
                    export_url = f"{GOOGLE_DRIVE_API_BASE}/files/{file_id}/export"
                    res = await client.get(
                        export_url,
                        headers=headers,
                        params={"mimeType": "application/vnd.openxmlformats-officedocument.wordprocessingml.document"},
                    )
                elif mime_type == "application/vnd.google-apps.spreadsheet":
                    # Export Google Sheet as XLSX
                    export_url = f"{GOOGLE_DRIVE_API_BASE}/files/{file_id}/export"
                    res = await client.get(
                        export_url,
                        headers=headers,
                        params={"mimeType": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"},
                    )
                elif mime_type == "application/vnd.google-apps.presentation":
                    # Export Google Slide as PPTX
                    export_url = f"{GOOGLE_DRIVE_API_BASE}/files/{file_id}/export"
                    res = await client.get(
                        export_url,
                        headers=headers,
                        params={"mimeType": "application/vnd.openxmlformats-officedocument.presentationml.presentation"},
                    )
                else:
                    # Standard binary or text download
                    download_url = f"{GOOGLE_DRIVE_API_BASE}/files/{file_id}"
                    res = await client.get(download_url, headers=headers, params={"alt": "media"})

                if res.status_code != 200:
                    raise ValueError(f"Failed to download Google Drive file: {res.text}")
                return res.content

    async def get_delta_changes(self, access_token: str, page_token: Optional[str] = None) -> Dict[str, Any]:
        """
        Queries Google Drive Changes API for incremental updates.
        """
        with trace_gdrive_call("get_delta_changes", "/changes"):
            headers = {"Authorization": f"Bearer {access_token}"}
            async with httpx.AsyncClient() as client:
                # 1. Fetch startPageToken if none provided
                if not page_token:
                    start_res = await client.get(
                        f"{GOOGLE_DRIVE_API_BASE}/changes/startPageToken",
                        headers=headers,
                    )
                    if start_res.status_code != 200:
                        raise ValueError(f"Failed to fetch Google startPageToken: {start_res.text}")
                    page_token = start_res.json().get("startPageToken")

                # 2. Query changes
                params = {
                    "pageToken": page_token,
                    "fields": "nextPageToken, newStartPageToken, changes(fileId, removed, file(id, name, mimeType, modifiedTime, webViewLink))",
                }
                res = await client.get(f"{GOOGLE_DRIVE_API_BASE}/changes", headers=headers, params=params)
                if res.status_code != 200:
                    raise ValueError(f"Failed to fetch Google Drive changes: {res.text}")
                return res.json()

    async def create_watch_channel(
        self,
        access_token: str,
        channel_id: str,
        webhook_url: str,
        token: str,
    ) -> Dict[str, Any]:
        """Subscribes to Google Drive push notification webhooks."""
        with trace_gdrive_call("create_watch_channel", "/changes/watch"):
            headers = {"Authorization": f"Bearer {access_token}"}
            body = {
                "id": channel_id,
                "type": "web_hook",
                "address": webhook_url,
                "token": token,
            }
            async with httpx.AsyncClient() as client:
                res = await client.post(f"{GOOGLE_DRIVE_API_BASE}/changes/watch", headers=headers, json=body)
                if res.status_code != 200:
                    raise ValueError(f"Failed to create Google Drive watch channel: {res.text}")
                return res.json()


# ==============================================================================
# Realistic Mock Google Drive Provider for Local Dev, Sandbox, and E2E Testing
# ==============================================================================
class MockGoogleDriveProvider:
    """
    Simulates Google Drive folders, Google Docs, Sheets, Slides, and binary files
    with dynamic, valid binary generation (PDF, DOCX, XLSX, PPTX, TXT) and caching.
    """
    def __init__(self):
        self._file_cache: Dict[str, bytes] = {}
        self._mock_tree = {
            "root": [
                {
                    "id": "gdrive_folder_corporate",
                    "name": "Corporate Shared Drive",
                    "is_folder": True,
                    "path": "/Corporate Shared Drive",
                    "drive_type": "google_drive",
                },
                {
                    "id": "gdrive_folder_tech",
                    "name": "Cloud & Engineering Specs",
                    "is_folder": True,
                    "path": "/Cloud & Engineering Specs",
                    "drive_type": "google_drive",
                },
                {
                    "id": "file_gdrive_welcome",
                    "name": "Google Drive System Overview.txt",
                    "is_folder": False,
                    "size": 1820,
                    "mime_type": "text/plain",
                    "created_date": "2026-03-01T10:00:00Z",
                    "modified_date": "2026-09-10T14:30:00Z",
                    "web_url": "https://drive.google.com/file/d/file_gdrive_welcome/view",
                    "path": "/Google Drive System Overview.txt",
                    "drive_type": "google_drive",
                },
            ],
            "gdrive_folder_corporate": [
                {
                    "id": "file_gdrive_doc",
                    "name": "Enterprise_AI_Security_Policy.gdoc",
                    "is_folder": False,
                    "size": 72400,
                    "mime_type": "application/vnd.google-apps.document",
                    "created_date": "2026-02-15T09:00:00Z",
                    "modified_date": "2026-09-15T11:45:00Z",
                    "web_url": "https://docs.google.com/document/d/file_gdrive_doc/edit",
                    "path": "/Corporate Shared Drive/Enterprise_AI_Security_Policy.gdoc",
                    "drive_type": "google_drive",
                },
                {
                    "id": "file_gdrive_sheet",
                    "name": "Global_Infrastructure_Costs_2026.gsheet",
                    "is_folder": False,
                    "size": 51200,
                    "mime_type": "application/vnd.google-apps.spreadsheet",
                    "created_date": "2026-01-20T08:15:00Z",
                    "modified_date": "2026-09-18T16:20:00Z",
                    "web_url": "https://docs.google.com/spreadsheets/d/file_gdrive_sheet/edit",
                    "path": "/Corporate Shared Drive/Global_Infrastructure_Costs_2026.gsheet",
                    "drive_type": "google_drive",
                },
                {
                    "id": "file_gdrive_slide",
                    "name": "Executive_Briefing_Q3.gslides",
                    "is_folder": False,
                    "size": 94800,
                    "mime_type": "application/vnd.google-apps.presentation",
                    "created_date": "2026-03-05T14:00:00Z",
                    "modified_date": "2026-09-20T10:10:00Z",
                    "web_url": "https://docs.google.com/presentation/d/file_gdrive_slide/edit",
                    "path": "/Corporate Shared Drive/Executive_Briefing_Q3.gslides",
                    "drive_type": "google_drive",
                },
            ],
            "gdrive_folder_tech": [
                {
                    "id": "file_gdrive_pdf",
                    "name": "Kubernetes_Hybrid_Cloud_Architecture.pdf",
                    "is_folder": False,
                    "size": 91200,
                    "mime_type": "application/pdf",
                    "created_date": "2026-02-10T12:00:00Z",
                    "modified_date": "2026-08-30T17:00:00Z",
                    "web_url": "https://drive.google.com/file/d/file_gdrive_pdf/view",
                    "path": "/Cloud & Engineering Specs/Kubernetes_Hybrid_Cloud_Architecture.pdf",
                    "drive_type": "google_drive",
                },
                {
                    "id": "file_gdrive_postmortem",
                    "name": "Incident_Postmortem_Database_Failover.txt",
                    "is_folder": False,
                    "size": 4250,
                    "mime_type": "text/plain",
                    "created_date": "2026-04-12T15:30:00Z",
                    "modified_date": "2026-07-22T09:40:00Z",
                    "web_url": "https://drive.google.com/file/d/file_gdrive_postmortem/view",
                    "path": "/Cloud & Engineering Specs/Incident_Postmortem_Database_Failover.txt",
                    "drive_type": "google_drive",
                },
            ],
        }

    async def list_drive_items(self, folder_id: str = "root") -> List[Dict[str, Any]]:
        with trace_gdrive_call("mock_list_drive_items", f"/folders/{folder_id}"):
            return self._mock_tree.get(folder_id, [])

    async def get_item_by_id(self, item_id: str) -> Optional[Dict[str, Any]]:
        for folder, items in self._mock_tree.items():
            for item in items:
                if item["id"] == item_id:
                    return item
        return None

    async def download_file_bytes(self, item_id: str, mime_type: str = "") -> bytes:
        with trace_gdrive_call("mock_download_file_bytes", f"/files/{item_id}"):
            if item_id in self._file_cache:
                return self._file_cache[item_id]

            if item_id == "file_gdrive_pdf":
                import fitz
                doc = fitz.open()
                page1 = doc.new_page()
                page1.insert_text(
                    (50, 72),
                    "KUBERNETES HYBRID CLOUD ARCHITECTURE (GOOGLE CLOUD + ONEDRIVE)\n\n"
                    "1. Multi-Cloud Ingress & Service Mesh:\n"
                    "Workloads span GKE (Google Kubernetes Engine) and Azure AKS with an Istio dual-control plane.\n"
                    "Workload identity federation links Google IAM roles to Azure Active Directory service principals.\n\n"
                    "2. RAG Microservices Topology:\n"
                    "- API Gateway router routes storage requests to OneDrive Graph API or Google Drive v3.\n"
                    "- Embeddings pipeline runs on GPUs with pgvector running inside a resilient high-availability cluster.\n"
                    "- Zero-trust mTLS secures all pod-to-pod communications."
                )
                page2 = doc.new_page()
                page2.insert_text(
                    (50, 72),
                    "Section 2: High Availability & Disaster Recovery\n\n"
                    "RPO (Recovery Point Objective): Under 60 seconds with asynchronous pgvector replication.\n"
                    "RTO (Recovery Time Objective): Under 5 minutes with automated DNS failover."
                )
                buf = io.BytesIO()
                doc.save(buf)
                doc.close()
                content = buf.getvalue()

            elif item_id == "file_gdrive_doc":
                import docx
                doc = docx.Document()
                doc.add_heading("Enterprise AI Security Policy 2026", level=0)
                doc.add_heading("Data Loss Prevention (DLP)", level=1)
                doc.add_paragraph("All queries executed through the RAG assistant are scrubbed for PII before prompt synthesis.")
                doc.add_paragraph("Documents stored in Google Shared Drives retain strict ACL enforcement mapping to user JWT tokens.")
                doc.add_heading("Model Access Guidelines", level=1)
                doc.add_paragraph("Only approved models (e.g. GPT-4o, Claude 3.5, Gemini 1.5) with zero data retention may process enterprise text.")
                buf = io.BytesIO()
                doc.save(buf)
                content = buf.getvalue()

            elif item_id == "file_gdrive_sheet":
                import openpyxl
                wb = openpyxl.Workbook()
                ws = wb.active
                ws.title = "Cloud Infrastructure 2026"
                ws.append(["Service Provider", "Category", "Q1 Actual ($)", "Q2 Forecast ($)", "Q3 Target ($)", "Annual Budget ($)"])
                ws.append(["Google Cloud Platform", "Kubernetes & BigQuery", 185000, 195000, 200000, 780000])
                ws.append(["Microsoft Azure", "Azure OpenAI & OneDrive Storage", 145000, 150000, 155000, 600000])
                ws.append(["Datadog & Langfuse", "Observability & Evals", 42000, 44000, 45000, 175000])
                ws.append(["Cloudflare", "CDN & WAF Security", 18000, 18000, 18000, 72000])
                buf = io.BytesIO()
                wb.save(buf)
                content = buf.getvalue()

            elif item_id == "file_gdrive_slide":
                from pptx import Presentation
                prs = Presentation()
                slide = prs.slides.add_slide(prs.slide_layouts[0])
                slide.shapes.title.text = "Executive Briefing: Q3 Enterprise AI"
                slide.shapes.placeholders[1].text = "Multi-Cloud RAG Platform (Google Drive + OneDrive)"
                slide2 = prs.slides.add_slide(prs.slide_layouts[1])
                slide2.shapes.title.text = "Strategic Milestones Achieved"
                slide2.shapes.placeholders[1].text = (
                    "• Bi-directional real-time sync with Google Drive v3 and Microsoft Graph.\n"
                    "• Sub-second semantic search over 100k+ enterprise document chunks with pgvector.\n"
                    "• Full observability coverage with OpenTelemetry distributed tracing and Langfuse evaluations."
                )
                buf = io.BytesIO()
                prs.save(buf)
                content = buf.getvalue()

            elif item_id == "file_gdrive_welcome":
                content = (
                    "GOOGLE DRIVE ENTERPRISE RAG CONNECTOR\n"
                    "====================================\n\n"
                    "This connector bridges Google Drive Shared Drives with our pgvector semantic indexing engine.\n"
                    "Supported File Formats:\n"
                    "- Google Docs (.gdoc) via native DOCX export\n"
                    "- Google Sheets (.gsheet) via native XLSX export\n"
                    "- Google Slides (.gslides) via native PPTX export\n"
                    "- Standard PDFs, Text files, and Markdown\n"
                    "Security: Per-tenant row-level access control enforced via user_id foreign keys.\n"
                ).encode("utf-8")

            elif item_id == "file_gdrive_postmortem":
                content = (
                    "INCIDENT POSTMORTEM: PostgreSQL Database Failover\n"
                    "Date: 2026-07-22 | Severity: SEV-2\n\n"
                    "Summary: Primary pgvector cluster experienced memory pressure during large batch re-indexing.\n"
                    "Root Cause: Chunk batch size was unbounded for large XLSX spreadsheets with >10,000 rows.\n"
                    "Corrective Action: Implemented structure-aware chunking with 1,000-character windows and 150-char overlap.\n"
                    "Resolution Time: 14 minutes total downtime. No customer data was lost."
                ).encode("utf-8")

            else:
                content = f"Mock Google Drive content for item {item_id}".encode("utf-8")

            self._file_cache[item_id] = content
            return content

    def set_mock_file_content(self, item_id: str, new_bytes: bytes):
        self._file_cache[item_id] = new_bytes

    def clear_cache(self):
        self._file_cache.clear()

    async def get_delta_changes(self, access_token: str = "", page_token: Optional[str] = None) -> Dict[str, Any]:
        with trace_gdrive_call("mock_get_delta_changes", "/changes"):
            items = []
            for folder_items in self._mock_tree.values():
                for it in folder_items:
                    items.append({
                        "fileId": it["id"],
                        "removed": False,
                        "file": {
                            "id": it["id"],
                            "name": it["name"],
                            "mimeType": it.get("mime_type", ""),
                            "modifiedTime": it.get("modified_date"),
                            "webViewLink": it.get("web_url"),
                        }
                    })
            return {
                "nextPageToken": None,
                "newStartPageToken": "mock_gdrive_start_page_token_999",
                "changes": items,
            }

    async def create_watch_channel(
        self,
        access_token: str = "",
        channel_id: str = "",
        webhook_url: str = "",
        token: str = "",
    ) -> Dict[str, Any]:
        with trace_gdrive_call("mock_create_watch_channel", "/changes/watch"):
            return {
                "kind": "api#channel",
                "id": channel_id or "mock_channel_uuid_123",
                "resourceId": "mock_resource_id_456",
                "resourceUri": "https://www.googleapis.com/drive/v3/changes",
                "expiration": str(int(datetime.now(timezone.utc).timestamp() * 1000 + 604800000)),
            }


# Service instances
google_drive_service = GoogleDriveService()
mock_google_drive_provider = MockGoogleDriveProvider()


def get_google_drive_service():
    """Returns the mock provider in development mode, or the real GoogleDriveService in production."""
    if settings.DEV_MOCK_GDRIVE:
        return mock_google_drive_provider
    return google_drive_service
