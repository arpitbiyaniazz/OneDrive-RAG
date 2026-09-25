import io
import logging
from datetime import datetime, timezone
from typing import Any, Dict, List, Optional
import httpx

from app.core.config import settings
from app.core.telemetry.otel_setup import trace_onedrive_call

logger = logging.getLogger(__name__)

GRAPH_BASE_URL = "https://graph.microsoft.com/v1.0"
LOGIN_BASE_URL = "https://login.microsoftonline.com"


class MicrosoftGraphService:
    def __init__(self):
        self.client_id = settings.MICROSOFT_CLIENT_ID
        self.client_secret = settings.MICROSOFT_CLIENT_SECRET
        self.tenant_id = settings.MICROSOFT_TENANT_ID
        self.scopes = ["User.Read", "Files.Read.All", "offline_access"]

    def get_auth_url(self, redirect_uri: str, state: str) -> str:
        """Constructs Microsoft OAuth 2.0 authorization URL."""
        scope_str = "%20".join(self.scopes)
        return (
            f"{LOGIN_BASE_URL}/{self.tenant_id}/oauth2/v2.0/authorize"
            f"?client_id={self.client_id}"
            f"&response_type=code"
            f"&redirect_uri={redirect_uri}"
            f"&response_mode=query"
            f"&scope={scope_str}"
            f"&state={state}"
        )

    async def exchange_code(self, code: str, redirect_uri: str) -> Dict[str, Any]:
        """Exchanges authorization code for access and refresh tokens."""
        token_url = f"{LOGIN_BASE_URL}/{self.tenant_id}/oauth2/v2.0/token"
        data = {
            "client_id": self.client_id,
            "client_secret": self.client_secret,
            "grant_type": "authorization_code",
            "code": code,
            "redirect_uri": redirect_uri,
            "scope": " ".join(self.scopes),
        }
        async with httpx.AsyncClient() as client:
            res = await client.post(token_url, data=data)
            if res.status_code != 200:
                logger.error(f"Failed to exchange Microsoft code: {res.text}")
                raise ValueError(f"Microsoft OAuth failed: {res.text}")
            return res.json()

    async def get_user_profile(self, access_token: str) -> Dict[str, Any]:
        """Fetches the user profile from Microsoft Graph /me."""
        with trace_onedrive_call("get_user_profile", "/me"):
            headers = {"Authorization": f"Bearer {access_token}"}
            async with httpx.AsyncClient() as client:
                res = await client.get(f"{GRAPH_BASE_URL}/me", headers=headers)
                if res.status_code != 200:
                    raise ValueError(f"Failed to fetch Microsoft profile: {res.text}")
                return res.json()

    async def list_drive_items(self, access_token: str, item_id: str = "root") -> List[Dict[str, Any]]:
        """Lists folders and files in a specific OneDrive folder."""
        with trace_onedrive_call("list_drive_items", f"/me/drive/items/{item_id}/children"):
            headers = {"Authorization": f"Bearer {access_token}"}
            url = f"{GRAPH_BASE_URL}/me/drive/{'root' if item_id == 'root' else f'items/{item_id}'}/children"
            async with httpx.AsyncClient() as client:
                res = await client.get(url, headers=headers)
                if res.status_code != 200:
                    raise ValueError(f"Failed to list OneDrive items: {res.text}")
                data = res.json()
                items = data.get("value", [])
                formatted = []
                for item in items:
                    is_folder = "folder" in item
                    formatted.append({
                        "id": item["id"],
                        "name": item["name"],
                        "is_folder": is_folder,
                        "size": item.get("size", 0),
                        "mime_type": item.get("file", {}).get("mimeType", "folder" if is_folder else "application/octet-stream"),
                        "created_date": item.get("createdDateTime"),
                        "modified_date": item.get("lastModifiedDateTime"),
                        "web_url": item.get("webUrl"),
                        "parent_id": item.get("parentReference", {}).get("id"),
                        "path": item.get("parentReference", {}).get("path", "").replace("/drive/root:", "") + "/" + item["name"],
                    })
                return formatted

    async def download_file_bytes(self, access_token: str, item_id: str) -> bytes:
        """Downloads file contents from OneDrive."""
        with trace_onedrive_call("download_file", f"/me/drive/items/{item_id}/content"):
            headers = {"Authorization": f"Bearer {access_token}"}
            url = f"{GRAPH_BASE_URL}/me/drive/items/{item_id}/content"
            async with httpx.AsyncClient(follow_redirects=True) as client:
                res = await client.get(url, headers=headers)
                if res.status_code != 200:
                    raise ValueError(f"Failed to download OneDrive file: {res.text}")
                return res.content


# ==============================================================================
# Realistic Mock OneDrive Provider for Local Development and Immediate Testing
# ==============================================================================
class MockOneDriveProvider:
    """
    Provides a realistic OneDrive structure with realistic sample documents
    (PDF, DOCX, XLSX, PPTX, TXT) across HR, Finance, and Engineering departments.
    """
    def __init__(self):
        self._file_cache: Dict[str, bytes] = {}
        self._mock_tree = {
            "root": [
                {"id": "folder_hr", "name": "HR", "is_folder": True, "path": "/HR"},
                {"id": "folder_finance", "name": "Finance", "is_folder": True, "path": "/Finance"},
                {"id": "folder_engineering", "name": "Engineering", "is_folder": True, "path": "/Engineering"},
                {
                    "id": "file_welcome",
                    "name": "Company Overview.txt",
                    "is_folder": False,
                    "size": 1420,
                    "mime_type": "text/plain",
                    "created_date": "2026-01-10T09:00:00Z",
                    "modified_date": "2026-08-15T14:20:00Z",
                    "web_url": "https://onedrive.live.com/view.aspx?id=file_welcome",
                    "path": "/Company Overview.txt",
                }
            ],
            "folder_hr": [
                {
                    "id": "file_hr_leave",
                    "name": "Leave Policy 2026.pdf",
                    "is_folder": False,
                    "size": 84500,
                    "mime_type": "application/pdf",
                    "created_date": "2026-02-01T10:00:00Z",
                    "modified_date": "2026-08-20T11:30:00Z",
                    "web_url": "https://onedrive.live.com/view.aspx?id=file_hr_leave",
                    "path": "/HR/Leave Policy 2026.pdf",
                },
                {
                    "id": "file_hr_handbook",
                    "name": "Employee Handbook.docx",
                    "is_folder": False,
                    "size": 65400,
                    "mime_type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    "created_date": "2026-01-15T12:00:00Z",
                    "modified_date": "2026-07-10T16:45:00Z",
                    "web_url": "https://onedrive.live.com/view.aspx?id=file_hr_handbook",
                    "path": "/HR/Employee Handbook.docx",
                }
            ],
            "folder_finance": [
                {
                    "id": "file_fin_budget",
                    "name": "Budget 2026.xlsx",
                    "is_folder": False,
                    "size": 42100,
                    "mime_type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    "created_date": "2026-01-05T08:30:00Z",
                    "modified_date": "2026-09-01T09:15:00Z",
                    "web_url": "https://onedrive.live.com/view.aspx?id=file_fin_budget",
                    "path": "/Finance/Budget 2026.xlsx",
                },
                {
                    "id": "file_fin_expense",
                    "name": "Expense Policy.txt",
                    "is_folder": False,
                    "size": 3120,
                    "mime_type": "text/plain",
                    "created_date": "2026-03-12T14:00:00Z",
                    "modified_date": "2026-06-18T10:00:00Z",
                    "web_url": "https://onedrive.live.com/view.aspx?id=file_fin_expense",
                    "path": "/Finance/Expense Policy.txt",
                }
            ],
            "folder_engineering": [
                {
                    "id": "file_eng_api",
                    "name": "API Documentation.pdf",
                    "is_folder": False,
                    "size": 128400,
                    "mime_type": "application/pdf",
                    "created_date": "2026-04-10T11:00:00Z",
                    "modified_date": "2026-08-25T15:30:00Z",
                    "web_url": "https://onedrive.live.com/view.aspx?id=file_eng_api",
                    "path": "/Engineering/API Documentation.pdf",
                },
                {
                    "id": "file_eng_arch",
                    "name": "Microservices Architecture.pptx",
                    "is_folder": False,
                    "size": 94200,
                    "mime_type": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
                    "created_date": "2026-05-02T13:00:00Z",
                    "modified_date": "2026-08-30T17:00:00Z",
                    "web_url": "https://onedrive.live.com/view.aspx?id=file_eng_arch",
                    "path": "/Engineering/Microservices Architecture.pptx",
                }
            ]
        }

    async def list_drive_items(self, item_id: str = "root") -> List[Dict[str, Any]]:
        with trace_onedrive_call("mock_list_items", item_id):
            return self._mock_tree.get(item_id, [])

    async def get_item_by_id(self, item_id: str) -> Optional[Dict[str, Any]]:
        for folder_items in self._mock_tree.values():
            for item in folder_items:
                if item["id"] == item_id:
                    return item
        return None

    async def download_file_bytes(self, item_id: str) -> bytes:
        with trace_onedrive_call("mock_download_file", item_id):
            if item_id in self._file_cache:
                return self._file_cache[item_id]

            content: bytes = b""
            # Generate realistic document content on the fly
            if item_id == "file_welcome":
                content = (
                    "Welcome to Contoso Enterprise Knowledge Base.\n"
                    "We are a forward-looking technology enterprise founded in 2020.\n"
                    "Our core departments include HR, Finance, and Engineering.\n"
                    "All policies are updated on OneDrive regularly."
                ).encode("utf-8")
            elif item_id == "file_fin_expense":
                content = (
                    "Contoso Expense Policy 2026:\n"
                    "1. Meals during business travel are reimbursed up to $75/day.\n"
                    "2. Flights over 4 hours may be booked in Premium Economy with manager approval.\n"
                    "3. All receipts must be uploaded within 30 days of expense occurrence."
                ).encode("utf-8")
            elif item_id == "file_hr_leave":
                import fitz
                doc = fitz.open()
                page = doc.new_page()
                page.insert_text(
                    (50, 72),
                    "CONTOSO ENTERPRISE LEAVE POLICY 2026\n\n"
                    "1. Annual Leave Allowance:\n"
                    "Employees are entitled to 18 annual leave days per calendar year.\n"
                    "Leave requests must be submitted through the portal at least two weeks in advance.\n\n"
                    "2. Sick Leave:\n"
                    "Employees receive 10 paid sick days annually. A medical certificate is required for absences\n"
                    "exceeding 3 consecutive business days.\n\n"
                    "3. Carry Over:\n"
                    "Up to 5 unused leave days can be carried over into the first quarter of the following year."
                )
                page2 = doc.new_page()
                page2.insert_text(
                    (50, 72),
                    "Page 2: Parental and Special Leave\n\n"
                    "Primary caregivers receive 16 weeks of fully paid parental leave.\n"
                    "Secondary caregivers receive 4 weeks of paid leave.\n"
                    "Bereavement leave provides up to 5 paid days for immediate family members."
                )
                buf = io.BytesIO()
                doc.save(buf)
                doc.close()
                content = buf.getvalue()
            elif item_id == "file_hr_handbook":
                import docx
                doc = docx.Document()
                doc.add_heading("Employee Handbook 2026", level=0)
                doc.add_heading("Leave Management", level=1)
                doc.add_paragraph("All employees must record their time off in the HR portal.")
                doc.add_paragraph("Annual leave accrues on a monthly basis at the rate of 1.5 days per month.")
                doc.add_heading("Remote Work Policy", level=1)
                doc.add_paragraph("Employees may work remotely up to 3 days per week with team alignment.")
                buf = io.BytesIO()
                doc.save(buf)
                content = buf.getvalue()
            elif item_id == "file_fin_budget":
                import openpyxl
                wb = openpyxl.Workbook()
                ws = wb.active
                ws.title = "Budget 2026"
                ws.append(["Department", "Q1 Budget", "Q2 Budget", "Q3 Budget", "Q4 Budget", "Total 2026"])
                ws.append(["Engineering", 1200000, 1350000, 1400000, 1500000, 5450000])
                ws.append(["Marketing", 450000, 500000, 550000, 600000, 2100000])
                ws.append(["Operations", 300000, 310000, 320000, 330000, 1260000])
                ws.append(["HR", 180000, 190000, 200000, 210000, 780000])
                buf = io.BytesIO()
                wb.save(buf)
                content = buf.getvalue()
            elif item_id == "file_eng_api":
                import fitz
                doc = fitz.open()
                page = doc.new_page()
                page.insert_text(
                    (50, 72),
                    "ENGINEERING API DOCUMENTATION & ARCHITECTURE\n\n"
                    "1. Authentication Mechanism:\n"
                    "All microservices utilize OAuth 2.0 with JWT Bearer tokens for API authentication.\n"
                    "Tokens are signed with RS256 and expire after 60 minutes.\n"
                    "API rate limits are enforced at 100 requests per minute per tenant.\n\n"
                    "2. Endpoints:\n"
                    "- GET /api/v1/health: Service status\n"
                    "- POST /api/v1/ingest: Document pipeline intake\n"
                    "- POST /api/v1/chat: Streaming RAG endpoint with Server-Sent Events (SSE)."
                )
                buf = io.BytesIO()
                doc.save(buf)
                doc.close()
                content = buf.getvalue()
            elif item_id == "file_eng_arch":
                from pptx import Presentation
                prs = Presentation()
                slide = prs.slides.add_slide(prs.slide_layouts[0])
                slide.shapes.title.text = "Microservices Architecture"
                slide.shapes.placeholders[1].text = "OneDrive RAG Platform • Antigravity Engineering"
                slide2 = prs.slides.add_slide(prs.slide_layouts[1])
                slide2.shapes.title.text = "Data Pipelines and Retrieval"
                slide2.shapes.placeholders[1].text = (
                    "• Ingestion service chunks documents by structural sections.\n"
                    "• pgvector stores 1536-dimensional embeddings for cosine similarity.\n"
                    "• Metadata filtering guarantees strict user multi-tenancy."
                )
                buf = io.BytesIO()
                prs.save(buf)
                content = buf.getvalue()
            else:
                content = f"Mock document content for {item_id}".encode("utf-8")

            self._file_cache[item_id] = content
            return content

    def set_mock_file_content(self, item_id: str, new_bytes: bytes):
        """Allows testing file modifications by overwriting cached mock file bytes."""
        self._file_cache[item_id] = new_bytes

    def clear_cache(self):
        """Clears the cached mock file bytes."""
        self._file_cache.clear()


# Singletons
graph_service = MicrosoftGraphService()
mock_onedrive_provider = MockOneDriveProvider()


def get_onedrive_service():
    """Returns real Microsoft Graph service or mock sandbox based on configuration."""
    if settings.DEV_MOCK_ONEDRIVE:
        return mock_onedrive_provider
    return graph_service
