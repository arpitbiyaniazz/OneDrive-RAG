"""Generates a multi-format test corpus with planted facts ("needles").

Every supported format is covered, with facts placed where naive parsers lose them (tables in the
middle of a document, speaker notes, second sheets, grouped shapes, headers/footers), plus files
that must be rejected instead of indexed (legacy Office, images, archives, empty, scanned PDF).
"""
import io
import zipfile
from dataclasses import dataclass, field
from datetime import date
from typing import List, Optional


@dataclass
class Needle:
    fact: str  # text that must appear verbatim (whitespace-normalized) in some chunk
    question: str  # natural question answered only by this fact
    answer: str  # short string the answer must contain


@dataclass
class CorpusFile:
    name: str
    data: bytes
    mime: str
    should_index: bool = True
    needles: List[Needle] = field(default_factory=list)
    reject_reason: Optional[str] = None  # substring expected in the UnsupportedDocumentError


def _pdf_report() -> bytes:
    import pymupdf

    doc = pymupdf.open()
    body = [
        "Quarterly Operations Report\n\nThe Zephyr project code is ZX-4471. It covers the warehouse "
        "automation rollout across three regions.\n\nThe rollout started in January and is staffed by "
        "a cross-functional team of logistics and software engineers.",
        "Regional Results\n\nThe northern region reached 97 percent uptime, the best in the company. "
        "Maintenance windows were moved to Sunday nights to reduce disruption for night-shift work-\n"
        "ers in the packing halls.",
        "Risks and Mitigations\n\nThe main risk is the supplier lead time for conveyor motors, currently "
        "eleven weeks. The mitigation is a second supplier in Rotterdam.",
    ]
    for index, text in enumerate(body, start=1):
        page = doc.new_page()
        page.insert_text((50, 40), "ACME Logistics - Internal", fontsize=8)
        page.insert_textbox(pymupdf.Rect(50, 70, 545, 400), text, fontsize=11)
        page.insert_text((50, 810), f"Confidential - page {index}", fontsize=8)
    # A real table with ruling lines so table detection finds it.
    page = doc.new_page()
    page.insert_text((50, 60), "Budget by region", fontsize=12)
    cells = [["Region", "Budget (EUR)", "Owner"], ["North", "1,250,000", "Ines Duarte"], ["South", "980,500", "Tomas Weber"]]
    x0, y0, w, h = 50, 80, 160, 24
    for r, row in enumerate(cells):
        for c, value in enumerate(row):
            rect = pymupdf.Rect(x0 + c * w, y0 + r * h, x0 + (c + 1) * w, y0 + (r + 1) * h)
            page.draw_rect(rect, color=(0, 0, 0), width=0.7)
            page.insert_textbox(rect + (4, 5, -4, -2), value, fontsize=10)
    doc.set_toc([[1, "Overview", 1], [1, "Regional Results", 2], [1, "Risks", 3], [1, "Budget", 4]])
    buf = io.BytesIO()
    doc.save(buf)
    doc.close()
    return buf.getvalue()


def _pdf_scanned() -> bytes:
    import pymupdf

    src = pymupdf.open()
    page = src.new_page()
    page.insert_text((50, 80), "This text only exists as pixels.", fontsize=14)
    pix = page.get_pixmap(dpi=72)
    out = pymupdf.open()
    target = out.new_page()
    target.insert_image(target.rect, stream=pix.tobytes("png"))
    buf = io.BytesIO()
    out.save(buf)
    return buf.getvalue()


def _docx_handbook() -> bytes:
    import docx

    d = docx.Document()
    d.sections[0].footer.paragraphs[0].text = "Policy owner: People Operations (contact: hr-ops@acme.example)"
    d.add_heading("Employee Handbook", level=0)
    d.add_heading("Leave", level=1)
    d.add_paragraph("Employees receive 26 days of annual leave per calendar year.")
    d.add_heading("Parental leave", level=2)
    d.add_paragraph("Primary caregivers receive 18 weeks of fully paid parental leave.")
    d.add_paragraph("Notify your manager at least 8 weeks before the start date.", style="List Bullet")
    d.add_heading("Equipment allowance", level=1)
    d.add_paragraph("The allowance table below is reviewed every January.")
    table = d.add_table(rows=3, cols=3)
    for r, row in enumerate([["Role", "Laptop budget", "Monitor"], ["Engineer", "2400 EUR", ""], ["Designer", "2800 EUR", "Yes"]]):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
    nested = table.cell(2, 2).add_table(rows=1, cols=2)
    nested.cell(0, 0).text = "Monitor model"
    nested.cell(0, 1).text = "UltraView 32Q"
    d.add_heading("Travel", level=1)
    d.add_paragraph("Train travel is preferred for trips under 500 km.")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _xlsx_budget() -> bytes:
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append(["FY2026 budget overview"])
    ws.append([])
    ws.append(["Department", "Owner", "Approved"])
    ws.append(["Engineering", "Priya Nair", date(2026, 1, 15)])
    ws.append(["Marketing", "Leo Martin", date(2026, 2, 3)])
    q = wb.create_sheet("Quarterly")
    q.append(["Department", "Q1", "Q2", "Q3", "Q4"])
    q.append(["Engineering", 1200000, 1350000, 1400000, 1500000])
    q.append(["Marketing", 450000, 500000, 731250, 600000])
    q.append(["Total", "=SUM(B2:B3)", "=SUM(C2:C3)", "=SUM(D2:D3)", "=SUM(E2:E3)"])
    wide = wb.create_sheet("Headcount")
    wide.append([f"Month {i}" for i in range(1, 31)])
    wide.append([str(40 + i) for i in range(30)])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _pptx_strategy() -> bytes:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    s1 = prs.slides.add_slide(prs.slide_layouts[0])
    s1.shapes.title.text = "Strategy 2027"
    s1.placeholders[1].text = "Board offsite deck"
    s2 = prs.slides.add_slide(prs.slide_layouts[1])
    s2.shapes.title.text = "Three pillars"
    s2.placeholders[1].text = "Customer intimacy\nOperational excellence\nPlatform reuse"
    s2.notes_slide.notes_text_frame.text = "The pilot launch date is 14 March 2027 in Lisbon."
    s3 = prs.slides.add_slide(prs.slide_layouts[5])
    s3.shapes.title.text = "Targets"
    rows, cols = 3, 2
    shape = s3.shapes.add_table(rows, cols, Inches(1), Inches(2), Inches(6), Inches(1.5))
    for r, row in enumerate([["Metric", "Target"], ["NPS", "62"], ["Churn", "4.5%"]]):
        for c, value in enumerate(row):
            shape.table.cell(r, c).text = value
    s4 = prs.slides.add_slide(prs.slide_layouts[6])
    group = s4.shapes.add_group_shape()
    box = group.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1))
    box.text_frame.text = "Grouped callout: the partner program is named Northstar."
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def build_corpus() -> List[CorpusFile]:
    files = [
        CorpusFile("Operations Report.pdf", _pdf_report(), "application/pdf", needles=[
            Needle("The Zephyr project code is ZX-4471.", "What is the Zephyr project code?", "ZX-4471"),
            Needle("night-shift workers in the packing halls", "Why were maintenance windows moved to Sunday nights?", "night-shift"),
            Needle("Rotterdam", "What is the mitigation for the conveyor motor lead time?", "Rotterdam"),
            Needle("1,250,000", "What is the North region budget?", "1,250,000"),
        ]),
        CorpusFile("Scanned Contract.pdf", _pdf_scanned(), "application/pdf", should_index=False, reject_reason="text layer"),
        CorpusFile("Employee Handbook.docx", _docx_handbook(),
                   "application/vnd.openxmlformats-officedocument.wordprocessingml.document", needles=[
            Needle("Primary caregivers receive 18 weeks of fully paid parental leave.", "How much parental leave do primary caregivers get?", "18 weeks"),
            Needle("2800 EUR", "What is the laptop budget for designers?", "2800"),
            Needle("UltraView 32Q", "Which monitor model do designers get?", "UltraView 32Q"),
            Needle("hr-ops@acme.example", "What is the contact address of the policy owner?", "hr-ops@acme.example"),
        ]),
        CorpusFile("Budget FY2026.xlsx", _xlsx_budget(),
                   "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", needles=[
            Needle("731250", "What is the Q3 budget for Marketing?", "731250"),
            Needle("Priya Nair", "Who owns the Engineering budget?", "Priya Nair"),
            Needle("2026-01-15", "When was the Engineering budget approved?", "2026-01-15"),
            Needle("=SUM(D2:D3)", "How is the Q3 total calculated?", "SUM"),
        ]),
        CorpusFile("Strategy 2027.pptx", _pptx_strategy(),
                   "application/vnd.openxmlformats-officedocument.presentationml.presentation", needles=[
            Needle("The pilot launch date is 14 March 2027 in Lisbon.", "When and where is the pilot launch?", "14 March 2027"),
            Needle("Operational excellence", "What are the three strategy pillars?", "Operational excellence"),
            Needle("4.5%", "What is the churn target?", "4.5%"),
            Needle("Northstar", "What is the partner program called?", "Northstar"),
        ]),
        CorpusFile("notes-utf8.txt", "Café opening hours: 07:30 to 18:00 on weekdays.\n".encode(), "text/plain", needles=[
            Needle("07:30 to 18:00", "What are the café opening hours?", "07:30"),
        ]),
        CorpusFile("notes-bom.txt", "﻿The VPN gateway is vpn2.acme.example.\n".encode(), "text/plain", needles=[
            Needle("vpn2.acme.example", "What is the VPN gateway address?", "vpn2.acme.example"),
        ]),
        CorpusFile("notes-cp1252.txt", "Le comité se réunit le jeudi à 9h; budget €12 500.\n".encode("cp1252"), "text/plain", needles=[
            Needle("€12 500", "Quel est le budget du comité ?", "12 500"),
        ]),
        CorpusFile("notes-utf16.txt", "Parking badge renewals happen every April.\n".encode("utf-16"), "text/plain", needles=[
            Needle("Parking badge renewals happen every April.", "When are parking badges renewed?", "April"),
        ]),
        CorpusFile("Runbook.md", (
            b"# Incident Runbook\n\nIntro text.\n\n## Escalation\n\nPage the on-call lead via the code word BLUEFIN.\n\n"
            b"```\n# not a heading\nrestart --all\n```\n"
        ), "text/markdown", needles=[
            Needle("BLUEFIN", "What code word is used to page the on-call lead?", "BLUEFIN"),
        ]),
        CorpusFile("vendors.csv", b'Vendor,Contact,Contract end\n"Globex, Inc.",Ana Silva,2027-06-30\nInitech,Raj Patel,2026-12-31\n', "text/csv", needles=[
            Needle("Ana Silva", "Who is the contact at Globex?", "Ana Silva"),
        ]),
        CorpusFile("Release Notes.html", b"<html><head><title>Release 4.2</title><style>p{}</style></head><body><h1>Release 4.2</h1>"
                   b"<p>Single sign-on now supports the Okta provider.</p><script>var x=1;</script></body></html>", "text/html", needles=[
            Needle("Single sign-on now supports the Okta provider.", "Which provider does single sign-on support in release 4.2?", "Okta"),
        ]),
        CorpusFile("Plan.gdoc", _docx_handbook(), "application/vnd.google-apps.document", needles=[
            Needle("UltraView 32Q", "Which monitor model is listed in the plan?", "UltraView 32Q"),
        ]),
        CorpusFile("misnamed-pdf.txt", _pdf_report(), "text/plain", needles=[
            Needle("ZX-4471", "What is the Zephyr project code?", "ZX-4471"),
        ]),
        CorpusFile("empty.txt", b"", "text/plain", should_index=False, reject_reason="empty"),
        CorpusFile("blank.txt", b"   \n\t\n", "text/plain", should_index=False, reject_reason="empty"),
        CorpusFile("legacy.doc", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 2048, "application/msword", should_index=False, reject_reason="Legacy"),
        CorpusFile("photo.png", b"\x89PNG\r\n\x1a\n" + b"\x00" * 512, "image/png", should_index=False, reject_reason="Images"),
        CorpusFile("archive.zip", _zip_bytes(), "application/zip", should_index=False, reject_reason="ZIP"),
        CorpusFile("binary.bin", bytes(range(256)) * 8, "application/octet-stream", should_index=False, reject_reason="Binary"),
    ]
    big = ("Section filler sentence about routine operations. " * 30 + "\n\n") * 120
    big += "The archive retention period is 7 years."
    files.append(CorpusFile("large-log.txt", big.encode(), "text/plain", needles=[
        Needle("The archive retention period is 7 years.", "What is the archive retention period?", "7 years"),
    ]))
    return files


def _zip_bytes() -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as zf:
        zf.writestr("readme.txt", "hello")
    return buf.getvalue()
