"""Real-world document structures that naive parsers lose (each case was reproduced in a review)."""
import io
import re
import zipfile
from datetime import date

import pytest

from app.processors import UnsupportedDocumentError, parse_document
from app.processors.text import decode_text
from app.services.chunker import chunk_document

W = 'xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"'


def _text(data: bytes, name: str) -> str:
    chunks = chunk_document(parse_document(data, name), name)
    return re.sub(r"\s+", " ", "\n".join(c.context + "\n" + c.content for c in chunks))


def _rezip(data: bytes, replace=None, add=None) -> bytes:
    src = zipfile.ZipFile(io.BytesIO(data))
    out = io.BytesIO()
    with zipfile.ZipFile(out, "w", zipfile.ZIP_DEFLATED) as z:
        for info in src.infolist():
            z.writestr(info.filename, (replace or {}).get(info.filename, src.read(info.filename)))
        for name, content in (add or {}).items():
            z.writestr(name, content)
    return out.getvalue()


# --------------------------------------------------------------------------- Word

def _docx_with(xml_blocks):
    import docx
    from docx.oxml import parse_xml

    d = docx.Document()
    d.add_paragraph("Body text.")
    sect = d.element.body[-1]
    for xml in xml_blocks:
        sect.addprevious(parse_xml(xml))
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def test_word_content_controls_tracked_changes_fields_and_text_boxes():
    wps = ('xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006" '
           'xmlns:v="urn:schemas-microsoft-com:vml"')
    data = _docx_with([
        f'<w:sdt {W}><w:sdtContent><w:p><w:r><w:t>Client reference CR-90817.</w:t></w:r></w:p></w:sdtContent></w:sdt>',
        f'<w:p {W}><w:r><w:t xml:space="preserve">Effective </w:t></w:r><w:sdt><w:sdtContent><w:r><w:t>1 March 2027</w:t></w:r></w:sdtContent></w:sdt></w:p>',
        f'<w:p {W}><w:r><w:t xml:space="preserve">Notice is </w:t></w:r><w:del w:id="1" w:author="a"><w:r><w:delText>30</w:delText></w:r></w:del>'
        f'<w:ins w:id="2" w:author="a"><w:r><w:t>90</w:t></w:r></w:ins><w:r><w:t xml:space="preserve"> days.</w:t></w:r></w:p>',
        f'<w:p {W}><w:fldSimple w:instr=" DOCPROPERTY Company "><w:r><w:t>Northwind Traders</w:t></w:r></w:fldSimple></w:p>',
        f'<w:p {W}><w:smartTag w:uri="urn:x" w:element="City"><w:r><w:t>Rotterdam</w:t></w:r></w:smartTag></w:p>',
        f'<w:p {W} {wps}><w:r><mc:AlternateContent><mc:Choice Requires="wps"><w:pict><v:shape><v:textbox><w:txbxContent>'
        f'<w:p><w:r><w:t>Key contact Dana Whitfield.</w:t></w:r></w:p></w:txbxContent></v:textbox></v:shape></w:pict></mc:Choice>'
        f'<mc:Fallback><w:pict><v:shape><v:textbox><w:txbxContent><w:p><w:r><w:t>Key contact Dana Whitfield.</w:t></w:r></w:p>'
        f'</w:txbxContent></v:textbox></v:shape></w:pict></mc:Fallback></mc:AlternateContent></w:r></w:p>',
        f'<w:tbl {W}><w:tr><w:tc><w:p><w:r><w:t>Applicant</w:t></w:r></w:p></w:tc><w:tc><w:p><w:r><w:t>Dept</w:t></w:r></w:p></w:tc></w:tr>'
        f'<w:sdt><w:sdtContent><w:tr><w:tc><w:sdt><w:sdtContent><w:p><w:r><w:t>Jordan Mbeki</w:t></w:r></w:p></w:sdtContent></w:sdt></w:tc>'
        f'<w:tc><w:p><w:r><w:t>Finance</w:t></w:r></w:p></w:tc></w:tr></w:sdtContent></w:sdt></w:tbl>',
    ])
    text = _text(data, "contract.docx")
    for fact in ("Client reference CR-90817.", "Effective 1 March 2027", "Notice is 90 days.", "Northwind Traders",
                 "Rotterdam", "| Jordan Mbeki | Finance |"):
        assert fact in text, fact
    assert "30" not in text.replace("CR-90817", "")  # deleted text stays out
    assert text.count("Key contact Dana Whitfield.") == 1  # Choice only, never Choice and Fallback


def test_word_numbering_titles_custom_headings_footnotes_and_headers():
    import docx
    from docx.enum.style import WD_STYLE_TYPE

    d = docx.Document()
    d.add_heading("Annual Report 2026", level=0)
    d.add_heading("Results", level=1)
    d.add_heading("Costs", level=2)  # heading with nothing under it
    d.add_heading("Outlook", level=2)
    d.add_paragraph("Flat growth.")
    d.add_paragraph("Payment terms", style="List Number")
    d.add_paragraph("Termination", style="List Number")
    style = d.styles.add_style("Clause Heading", WD_STYLE_TYPE.PARAGRAPH)
    style.base_style = d.styles["Heading 1"]
    d.add_paragraph("Confidentiality", style="Clause Heading")
    d.add_paragraph("Secret for 5 years.")
    section = d.sections[0]
    section.different_first_page_header_footer = True
    section.first_page_header.paragraphs[0].text = "Document ID: DOC-2231"
    table = section.header.add_table(rows=1, cols=2, width=docx.shared.Inches(6))
    table.cell(0, 0).text, table.cell(0, 1).text = "Project", "Falcon"
    p = d.add_paragraph("Liability is capped")
    p._p.append(docx.oxml.parse_xml(f'<w:r {W}><w:footnoteReference w:id="1"/></w:r>'))
    buf = io.BytesIO()
    d.save(buf)
    footnotes = (f'<?xml version="1.0" encoding="UTF-8"?><w:footnotes {W}><w:footnote w:type="separator" w:id="-1">'
                 '<w:p><w:r><w:separator/></w:r></w:p></w:footnote><w:footnote w:id="1"><w:p><w:r>'
                 '<w:t>at 2x the annual fees.</w:t></w:r></w:p></w:footnote></w:footnotes>')
    z = zipfile.ZipFile(io.BytesIO(buf.getvalue()))
    types = z.read("[Content_Types].xml").decode().replace("</Types>", '<Override PartName="/word/footnotes.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"/></Types>')
    rels = z.read("word/_rels/document.xml.rels").decode().replace("</Relationships>", '<Relationship Id="rIdFn" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/footnotes" Target="footnotes.xml"/></Relationships>')
    data = _rezip(buf.getvalue(), replace={"[Content_Types].xml": types, "word/_rels/document.xml.rels": rels},
                  add={"word/footnotes.xml": footnotes})
    text = _text(data, "report.docx")
    for fact in ("Annual Report 2026", "Costs", "1. Payment terms", "2. Termination",
                 "Liability is capped (Footnote: at 2x the annual fees.)", "Document ID: DOC-2231", "Project | Falcon"):
        assert fact in text, fact
    chunks = chunk_document(parse_document(data, "report.docx"), "report.docx")
    assert any(c.section == "Confidentiality" and "Secret for 5 years." in c.content for c in chunks)  # custom heading style


def test_word_table_captions_vertical_merges_and_links():
    data = _docx_with([
        f'<w:tbl {W}><w:tblGrid><w:gridCol/><w:gridCol/><w:gridCol/></w:tblGrid>'
        f'<w:tr><w:tc><w:tcPr><w:gridSpan w:val="3"/></w:tcPr><w:p><w:r><w:t>Table 4: Professional fees</w:t></w:r></w:p></w:tc></w:tr>'
        f'<w:tr><w:tc><w:p><w:r><w:t>Site</w:t></w:r></w:p></w:tc><w:tc><w:p><w:r><w:t>System</w:t></w:r></w:p></w:tc><w:tc><w:p><w:r><w:t>Owner</w:t></w:r></w:p></w:tc></w:tr>'
        f'<w:tr><w:tc><w:tcPr><w:vMerge w:val="restart"/></w:tcPr><w:p><w:r><w:t>Dublin</w:t></w:r></w:p></w:tc><w:tc><w:p><w:r><w:t>Payroll</w:t></w:r></w:p></w:tc><w:tc><w:p><w:r><w:t>Aoife</w:t></w:r></w:p></w:tc></w:tr>'
        f'<w:tr><w:tc><w:tcPr><w:vMerge/></w:tcPr><w:p/></w:tc><w:tc><w:p><w:r><w:t>CRM</w:t></w:r></w:p></w:tc><w:tc><w:p><w:r><w:t>Liam</w:t></w:r></w:p></w:tc></w:tr></w:tbl>',
    ])
    parsed = parse_document(data, "systems.docx")
    tables = [e.text for e in parsed.elements if e.kind == "table"]
    assert tables and tables[0].startswith("| Site | System | Owner |") and "| Dublin | CRM | Liam |" in tables[0]
    assert any(e.text == "Table 4: Professional fees" for e in parsed.elements)


# --------------------------------------------------------------------------- Excel

def _xlsx(build) -> bytes:
    import openpyxl

    wb = openpyxl.Workbook()
    build(wb)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_excel_chart_sheets_merged_two_row_headers_and_stacked_tables():
    from openpyxl.chart import BarChart, Reference

    def build(wb):
        ws = wb.active
        ws.title = "Results"
        ws.append(["Region", "Office", "2025", None, "2026", None])
        ws.append([None, None, "Plan", "Actual", "Plan", "Actual"])
        for rng in ("C1:D1", "E1:F1", "A1:A2", "B1:B2"):
            ws.merge_cells(rng)
        ws.append(["North", "Oslo", 100, 90, 110, 125])
        ws.append([None, "Bergen", 80, 85, 90, 70])
        ws.merge_cells("A3:A4")
        ws.append([])
        ws.append(["Escalation level", "Response time", "Channel"])
        ws.append(["P1", "15 minutes", "Phone"])
        chart = BarChart()
        chart.add_data(Reference(ws, min_col=3, min_row=1, max_row=4))
        wb.create_chartsheet("Chart").add_chart(chart)

    text = _text(_xlsx(build), "results.xlsx")
    assert "| Region | Office | 2025 Plan | 2025 Actual | 2026 Plan | 2026 Actual |" in text
    assert "| North | Bergen | 80 | 85 | 90 | 70 |" in text  # merged region label filled down
    assert "| Escalation level | Response time | Channel |" in text and "| P1 | 15 minutes | Phone |" in text


def test_excel_preambles_year_headers_formats_hidden_sheets_and_profiles():
    def build(wb):
        ws = wb.active
        ws.title = "Export"
        ws.append(["Report:", "Open purchase orders"])
        ws.append([])
        ws.append(["Metric", 2024, 2025])
        ws.append(["Revenue", 41.2, 44.9])
        k = wb.create_sheet("KPIs")
        k.append(["KPI", "Value"])
        for label, value, fmt in [("Churn", 0.045, "0.0%"), ("Zip", 2134, "00000"), ("Hours", 1.5, "[h]:mm"),
                                  ("Cost", 1149.9999999999998, "#,##0.00")]:
            k.append([label, value])
            k.cell(row=k.max_row, column=2).number_format = fmt
        h = wb.create_sheet("Old")
        h.sheet_state = "hidden"
        h.append(["Item", "Price"])
        for i in range(6):
            h.append([f"Widget {i}", 10 * (i + 1)])

    parsed = parse_document(_xlsx(build), "kpi.xlsx")
    text = "\n".join(e.text for e in parsed.elements)
    assert "Report: Open purchase orders" in text and "| Metric | 2024 | 2025 |" in text
    for value in ("4.5%", "02134", "36:00", "1,150.00"):
        assert f"| {value} |" in text, value
    hidden = [e for e in parsed.elements if e.section == "Sheet: Old (hidden)"]
    assert hidden and any("6 data rows" in e.text and "Price: total 210" in e.text for e in hidden)


def test_excel_wrong_stored_dimension_does_not_hide_rows():
    def build(wb):
        ws = wb.active
        ws.append(["Order", "Amount"])
        ws.append(["SO-103", 1000])

    data = _xlsx(build)
    xml = zipfile.ZipFile(io.BytesIO(data)).read("xl/worksheets/sheet1.xml").decode()
    data = _rezip(data, replace={"xl/worksheets/sheet1.xml": re.sub(r"<dimension[^>]*/>", '<dimension ref="A1"/>', xml)})
    assert "| SO-103 | 1000 |" in _text(data, "orders.xlsx")


# --------------------------------------------------------------------------- PDF

def _pdf(build) -> bytes:
    import pymupdf

    doc = pymupdf.open()
    build(doc, pymupdf)
    return doc.tobytes()


def test_pdf_columns_hyphenation_and_captions():
    def build(doc, pymupdf):
        page = doc.new_page()
        for i, text in enumerate(["L1 warranty covers parts", "L2 for three years."]):
            page.insert_textbox(pymupdf.Rect(50, 100 + 60 * i, 280, 150 + 60 * i), text, fontsize=10)
        for i, text in enumerate(["R1 shipping is free", "R2 above 500 EUR."]):
            page.insert_textbox(pymupdf.Rect(310, 95 + 60 * i, 545, 145 + 60 * i), text, fontsize=10)
        page = doc.new_page()
        page.insert_text((50, 300), "Runs 2024-", fontsize=10)
        page.insert_text((50, 312), "2027 with infor-", fontsize=10)
        page.insert_text((50, 324), "mation on state-of-the-", fontsize=10)
        page.insert_text((50, 336), "art parts.", fontsize=10)

    first, second = parse_document(_pdf(build), "terms.pdf").elements[:2]
    assert [line[:2] for line in first.text.splitlines()[:4]] == ["L1", "L2", "R1", "R2"]
    assert "2024-\n2027" in second.text and "information" in second.text and "state-of-the-art" in second.text


def test_pdf_caption_touching_a_table_is_kept():
    def build(doc, pymupdf):
        page = doc.new_page()
        page.insert_text((40, 98), "Table 3: Penalty rates agreed with the regulator", fontsize=10)
        cells = [["Breach", "Penalty"], ["Late filing", "500 EUR"], ["Missing data", "1,200 EUR"]]
        for r, row in enumerate(cells):
            for c, value in enumerate(row):
                rect = pymupdf.Rect(40 + c * 175, 99.5 + r * 24, 40 + (c + 1) * 175, 99.5 + (r + 1) * 24)
                page.draw_rect(rect, color=(0, 0, 0), width=0.7)
                page.insert_textbox(rect + (4, 5, -4, -2), value, fontsize=10)

    text = _text(_pdf(build), "rates.pdf")
    assert "Table 3: Penalty rates agreed with the regulator" in text and "| Late filing | 500 EUR |" in text


# --------------------------------------------------------------------------- PowerPoint

def test_powerpoint_charts_alt_text_and_soft_breaks():
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.enum.chart import XL_CHART_TYPE
    from pptx.util import Inches

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Bookings"
    box = slide.shapes.add_textbox(Inches(1), Inches(1.5), Inches(4), Inches(1)).text_frame
    box.text = "Revenue grew\vby 20 percent"
    chart_data = CategoryChartData()
    chart_data.categories = ["Q1", "Q2"]
    chart_data.add_series("Bookings (kEUR)", (410, 455))
    slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(3), Inches(4), Inches(3), chart_data)
    picture = slide.shapes.add_picture(io.BytesIO(_png()), Inches(6), Inches(1))
    picture._element.xpath(".//p:cNvPr")[0].set("descr", "Org chart of the sales team")
    buf = io.BytesIO()
    prs.save(buf)
    text = _text(buf.getvalue(), "deck.pptx")
    for fact in ("Revenue grew by 20 percent", "| Category | Bookings (kEUR) |", "| Q2 | 455 |", "Image: Org chart of the sales team"):
        assert fact in text, fact


def _png() -> bytes:
    import pymupdf

    doc = pymupdf.open()
    return doc.new_page(width=40, height=40).get_pixmap().tobytes("png")


# --------------------------------------------------------------------------- text, Markdown, HTML, packages

@pytest.mark.parametrize("text,encoding", [
    ("Отпуск сотрудника: 28 календарных дней.", "cp1251"),
    ("有給休暇は年間二十日です。", "shift_jis"),
    ("Urlop wynosi 26 dni, zgłoszenie do kierownika.", "iso-8859-2"),
    ("Le comité se réunit le jeudi à 9h.", "cp1252"),
    ("Prénom: Zoë; Straße 12; Ελλάδα", "utf-16-le"),
])
def test_short_legacy_encodings_decode(text, encoding):
    assert decode_text(text.encode(encoding)) == text


def test_markdown_fences_setext_headings_and_tables():
    md = ("Project Atlas\n=============\n\nIntro.\n\n~~~bash\n# not a heading\n~~~\n\n## Pricing\n\n"
          "| Plan | Price |\n|---|---|\n" + "\n".join(f"| Plan-{i} | {i * 49} EUR |" for i in range(1, 80)) + "\n")
    chunks = chunk_document(parse_document(md.encode(), "atlas.md"), "atlas.md")
    assert all(c.section in ("Project Atlas", "Project Atlas > Pricing") for c in chunks)
    table_chunks = [c for c in chunks if "Plan-" in c.content]
    assert len(table_chunks) > 1 and all(c.content.startswith("| Plan | Price |") for c in table_chunks)


def test_html_definitions_titles_links_and_tables():
    html = ("<html><head><title>HR FAQ</title></head><body><svg><title>icon</title></svg><h1>Benefits</h1>"
            "<dl><dt>Gym allowance</dt><dd>45 EUR per month</dd></dl>"
            '<p>Submit via <a href="https://expenses.contoso.example/submit">the portal</a>.</p>'
            "<table><tr><th>Level</th><th>Days</th></tr><tr><td>Junior</td><td>25</td></tr></table></body></html>")
    parsed = parse_document(html.encode(), "faq.html")
    text = "\n".join(e.text for e in parsed.elements)
    assert "Gym allowance: 45 EUR per month" in text and "(https://expenses.contoso.example/submit)" in text
    assert "| Junior | 25 |" in text and {e.section for e in parsed.elements} == {"Benefits"}


def test_tab_separated_text_exports_are_tables():
    rows = ["Employee\tDepartment\tSalary"] + [f"Person {i}\tDept {i % 7}\t{50000 + i}" for i in range(150)]
    parsed = parse_document("\n".join(rows).encode("utf-16"), "salaries.txt")
    assert parsed.file_type == "csv"
    assert all(e.text.startswith("| Employee | Department | Salary |") for e in parsed.elements if e.kind == "table")


def test_macro_enabled_and_template_office_files_are_read():
    import docx

    d = docx.Document()
    d.add_paragraph("Retention is seven years.")
    buf = io.BytesIO()
    d.save(buf)
    types = zipfile.ZipFile(io.BytesIO(buf.getvalue())).read("[Content_Types].xml").decode().replace(
        "wordprocessingml.document.main+xml", "wordprocessingml.template.main+xml")
    data = _rezip(buf.getvalue(), replace={"[Content_Types].xml": types})
    assert "Retention is seven years." in parse_document(data, "policy.dotx").elements[0].text


def test_unreadable_files_do_not_leak_internal_details():
    with pytest.raises(UnsupportedDocumentError) as exc:
        parse_document(b"PK\x03\x04" + b"\0" * 64, "broken.docx")
    assert "<" not in exc.value.reason and "object at" not in exc.value.reason


def test_dates_render_as_iso():
    from app.processors.xlsx import format_cell

    assert format_cell(date(2026, 1, 15)) == "2026-01-15" and format_cell(0.30000000000000004) == "0.3"
